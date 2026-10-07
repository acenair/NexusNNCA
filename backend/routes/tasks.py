"""Routes / tasks — extracted from server.py without behavior changes."""
from fastapi import APIRouter
from pydantic import BaseModel
from fastapi import Cookie
from fastapi import File
from fastapi import HTTPException
from fastapi import Header
from typing import Optional
from fastapi import UploadFile
from datetime import datetime
from core import db
from core import get_current_user
from core import log_activity
from datetime import timedelta
from datetime import timezone
import uuid

router = APIRouter(prefix="/api")


@router.get("/tasks")
async def get_tasks(authorization: str = Header(None), session_token: str = Cookie(None), filter: Optional[str] = None):
    user = await get_current_user(authorization, session_token)
    
    query = {}
    if filter == "my":
        query["assigned_to"] = user["user_id"]
    elif filter == "overdue":
        query["status"] = {"$ne": "Completed"}
        query["due_date"] = {"$lt": datetime.now(timezone.utc).isoformat()}
    elif filter == "week":
        week_end = (datetime.now(timezone.utc) + timedelta(days=7)).isoformat()
        query["status"] = {"$ne": "Completed"}
        query["due_date"] = {"$lte": week_end}
    
    tasks = await db.tasks.find(query, {"_id": 0}).sort("due_date", 1).to_list(1000)
    return tasks


class CreateTaskRequest(BaseModel):
    title: str
    service_module: str
    due_date: str
    priority: str
    description: Optional[str] = None
    client_id: Optional[str] = None
    client_name: Optional[str] = None
    assigned_to: Optional[str] = None
    assigned_to_name: Optional[str] = None


@router.post("/tasks")
async def create_task(req: CreateTaskRequest,
                     authorization: str = Header(None), session_token: str = Cookie(None)):
    user = await get_current_user(authorization, session_token)
    
    task_id = f"task_{uuid.uuid4().hex[:12]}"
    task_doc = {
        "task_id": task_id,
        "title": req.title,
        "description": req.description,
        "service_module": req.service_module,
        "client_id": req.client_id,
        "client_name": req.client_name,
        "due_date": req.due_date,
        "priority": req.priority,
        "assigned_to": req.assigned_to or user["user_id"],
        "assigned_to_name": req.assigned_to_name,
        "status": "Pending",
        "created_by": user["user_id"],
        "created_by_name": user.get("name"),
        "created_at": datetime.now(timezone.utc).isoformat()
    }
    await db.tasks.insert_one(task_doc)
    await log_activity("Task created", f"Created task: {req.title}", user["user_id"], req.client_id, req.client_name)
    
    task_doc.pop("_id", None)
    return task_doc


@router.patch("/tasks/{task_id}")
async def update_task(task_id: str, status: Optional[str] = None, authorization: str = Header(None), session_token: str = Cookie(None)):
    user = await get_current_user(authorization, session_token)
    
    update_data = {}
    if status:
        update_data["status"] = status
        if status == "Completed":
            update_data["completed_at"] = datetime.now(timezone.utc).isoformat()
    
    await db.tasks.update_one({"task_id": task_id}, {"$set": update_data})
    return {"message": "Task updated"}


TASK_HEADER_MAP = {
    # title mappings
    "title": "title", "task title": "title", "task name": "title", "task": "title",
    "name": "title", "subject": "title", "description_title": "title",
    # description mappings
    "description": "description", "desc": "description", "details": "description",
    "notes": "description", "note": "description", "remarks": "description",
    # service module mappings
    "service module": "service_module", "service_module": "service_module",
    "service": "service_module", "module": "service_module", "category": "service_module",
    "department": "service_module", "type": "service_module",
    # client mappings
    "client": "client_name", "client name": "client_name", "client_name": "client_name",
    "company": "client_name", "entity": "client_name", "customer": "client_name",
    # assigned to mappings
    "assigned to": "assigned_to_name", "assigned_to": "assigned_to_name",
    "assignee": "assigned_to_name", "assigned": "assigned_to_name",
    "owner": "assigned_to_name", "responsible": "assigned_to_name",
    "staff": "assigned_to_name", "team member": "assigned_to_name",
    # due date mappings
    "due date": "due_date", "due_date": "due_date", "deadline": "due_date",
    "due": "due_date", "target date": "due_date", "end date": "due_date",
    # priority mappings
    "priority": "priority", "urgency": "priority", "importance": "priority",
    # status mappings
    "status": "status", "state": "status", "progress": "status",
}


def parse_date_flexible(val):
    """Try to parse various date formats into YYYY-MM-DD."""
    if not val:
        return None
    val = str(val).strip()
    for fmt in ["%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y", "%d %b %Y", "%d %B %Y", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d %H:%M:%S"]:
        try:
            return datetime.strptime(val, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    # Try pandas-like parsing
    try:
        from dateutil import parser as dateparser
        return dateparser.parse(val).strftime("%Y-%m-%d")
    except Exception:
        pass
    return val


@router.post("/tasks/bulk-upload")
async def bulk_upload_tasks(file: UploadFile = File(...), authorization: str = Header(None), session_token: str = Cookie(None)):
    """Parse CSV or XLSX file, auto-map columns, and create tasks in bulk."""
    user = await get_current_user(authorization, session_token)
    if user.get("role") != "partner":
        raise HTTPException(status_code=403, detail="Only partners can bulk upload tasks")

    # Read file content
    content = await file.read()
    filename = file.filename or ""

    rows = []
    headers = []

    if filename.lower().endswith(".xlsx") or filename.lower().endswith(".xls"):
        import openpyxl
        import io
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        ws = wb.active
        all_rows = list(ws.iter_rows(values_only=True))
        if len(all_rows) < 2:
            raise HTTPException(status_code=400, detail="File has no data rows")
        headers = [str(h).strip() if h else "" for h in all_rows[0]]
        for row in all_rows[1:]:
            rows.append([str(cell) if cell is not None else "" for cell in row])
    elif filename.lower().endswith(".csv"):
        import csv
        import io
        text = content.decode("utf-8-sig")
        reader = csv.reader(io.StringIO(text))
        all_rows = list(reader)
        if len(all_rows) < 2:
            raise HTTPException(status_code=400, detail="File has no data rows")
        headers = [h.strip() for h in all_rows[0]]
        rows = all_rows[1:]
    else:
        raise HTTPException(status_code=400, detail="Unsupported file format. Please upload CSV or XLSX.")

    # Auto-map headers
    column_map = {}
    for idx, h in enumerate(headers):
        normalized = h.lower().strip()
        if normalized in TASK_HEADER_MAP:
            column_map[idx] = TASK_HEADER_MAP[normalized]

    if "title" not in column_map.values():
        raise HTTPException(status_code=400, detail=f"Could not find a 'Title' column. Found headers: {headers}")

    # Resolve staff names to emails
    all_staff = await db.users.find({}, {"_id": 0, "name": 1, "email": 1}).to_list(100)
    staff_name_to_email = {u["name"].lower(): u["email"] for u in all_staff}

    # Resolve client names to IDs
    all_clients = await db.clients.find({}, {"_id": 0, "client_id": 1, "name": 1}).to_list(500)
    client_name_to_id = {c["name"].lower(): c["client_id"] for c in all_clients}

    created = []
    errors = []
    for row_idx, row in enumerate(rows, start=2):
        try:
            task_data = {}
            for col_idx, field in column_map.items():
                if col_idx < len(row):
                    task_data[field] = row[col_idx].strip()

            title = task_data.get("title", "").strip()
            if not title:
                continue  # Skip empty rows

            # Parse due date
            due_date = parse_date_flexible(task_data.get("due_date"))
            if not due_date:
                due_date = (datetime.now(timezone.utc) + timedelta(days=7)).strftime("%Y-%m-%d")

            # Resolve staff
            assignee_name = task_data.get("assigned_to_name", "").strip()
            assignee_email = staff_name_to_email.get(assignee_name.lower(), "")

            # Resolve client
            client_name = task_data.get("client_name", "").strip()
            client_id = client_name_to_id.get(client_name.lower(), "")

            # Normalize priority
            raw_priority = task_data.get("priority", "Medium").strip().capitalize()
            priority = raw_priority if raw_priority in ("High", "Medium", "Low") else "Medium"

            # Normalize status
            raw_status = task_data.get("status", "Pending").strip()
            status = raw_status if raw_status in ("Pending", "In Progress", "Completed") else "Pending"

            service_module = task_data.get("service_module", "General").strip() or "General"
            description = task_data.get("description", "").strip()

            task_id = f"task_{uuid.uuid4().hex[:12]}"
            task_doc = {
                "task_id": task_id,
                "title": title,
                "description": description or f"Bulk imported from {filename}",
                "service_module": service_module,
                "client_id": client_id,
                "client_name": client_name,
                "due_date": due_date,
                "priority": priority,
                "assigned_to": assignee_email or user.get("email"),
                "assigned_to_name": assignee_name or user.get("name"),
                "status": status,
                "created_by": user["user_id"],
                "created_by_name": user.get("name"),
                "created_at": datetime.now(timezone.utc).isoformat(),
                "source": "bulk_upload",
            }
            await db.tasks.insert_one(task_doc)
            created.append({"row": row_idx, "task_id": task_id, "title": title})
        except Exception as e:
            errors.append({"row": row_idx, "error": str(e)})

    await log_activity("Bulk task upload", f"Uploaded {len(created)} tasks from {filename}", user["user_id"])

    return {
        "total_rows": len(rows),
        "created": len(created),
        "errors": len(errors),
        "error_details": errors[:20],
        "mapped_columns": {headers[idx]: field for idx, field in column_map.items()},
        "message": f"Successfully imported {len(created)} task(s) from {filename}",
    }


@router.post("/tasks/bulk-preview")
async def preview_bulk_upload(file: UploadFile = File(...), authorization: str = Header(None), session_token: str = Cookie(None)):
    """Preview column mapping and first few rows without importing."""
    user = await get_current_user(authorization, session_token)
    if user.get("role") != "partner":
        raise HTTPException(status_code=403, detail="Only partners can bulk upload tasks")

    content = await file.read()
    filename = file.filename or ""

    rows = []
    headers = []

    if filename.lower().endswith(".xlsx") or filename.lower().endswith(".xls"):
        import openpyxl
        import io
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        ws = wb.active
        all_rows = list(ws.iter_rows(values_only=True))
        if len(all_rows) < 1:
            raise HTTPException(status_code=400, detail="File is empty")
        headers = [str(h).strip() if h else "" for h in all_rows[0]]
        for row in all_rows[1:6]:
            rows.append([str(cell) if cell is not None else "" for cell in row])
    elif filename.lower().endswith(".csv"):
        import csv
        import io
        text = content.decode("utf-8-sig")
        reader = csv.reader(io.StringIO(text))
        all_rows = list(reader)
        if len(all_rows) < 1:
            raise HTTPException(status_code=400, detail="File is empty")
        headers = [h.strip() for h in all_rows[0]]
        rows = all_rows[1:6]
    else:
        raise HTTPException(status_code=400, detail="Unsupported file format. Please upload CSV or XLSX.")

    column_map = {}
    unmapped = []
    for idx, h in enumerate(headers):
        normalized = h.lower().strip()
        if normalized in TASK_HEADER_MAP:
            column_map[h] = TASK_HEADER_MAP[normalized]
        else:
            unmapped.append(h)

    return {
        "filename": filename,
        "total_rows": len(all_rows) - 1 if filename.lower().endswith(".csv") else len(list(rows)),
        "headers": headers,
        "mapped_columns": column_map,
        "unmapped_columns": unmapped,
        "preview_rows": rows,
    }
