"""Digital Audit Workbook: template parsing, progress calc, Excel/PDF export."""
import io
import re
from datetime import datetime, timezone
from pathlib import Path

import xlrd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

TEMPLATE_ID = "std_statutory_audit_v2"
XLS_PATH = Path(__file__).parent / "seed_data" / "audit_index.xls"
# Sections that hold static/structural client data eligible for year-on-year roll-forward
ROLL_FORWARD_SECTION_ORDERS = {0, 1}  # Cover Page, Legal Documents


def parse_template():
    wb = xlrd.open_workbook(XLS_PATH)
    sections = []
    for order, sheet in enumerate(wb.sheets()):
        rows = [[str(sheet.cell_value(r, c)).strip() for c in range(sheet.ncols)] for r in range(sheet.nrows)]
        header_idx = next((i for i, row in enumerate(rows) if row and row[0] == "Sl No"), None)
        if header_idx is None:
            # Cover Page style sheets have no "Sl No" header row; data starts from the top
            header_idx = -1
        title = None
        for i in range(header_idx - 1, -1, -1):
            candidate = next((c for c in rows[i] if c), None)
            if candidate:
                title = candidate
                break
        if not title:
            title = re.sub(r"^\d+\)", "", sheet.name).strip()

        section_id = f"sec_{order + 1}"
        questions = []
        last_code = None
        blank_counters = {}
        used_codes = set()

        def _unique(candidate):
            if candidate not in used_codes:
                used_codes.add(candidate)
                return candidate
            i = 2
            while f"{candidate}_{i}" in used_codes:
                i += 1
            final = f"{candidate}_{i}"
            used_codes.add(final)
            return final

        for row in rows[header_idx + 1:]:
            code = row[0] if len(row) > 0 else ""
            text = row[1] if len(row) > 1 else ""
            remarks = row[2] if len(row) > 2 else ""
            if not text or (len(text) <= 5 and not code):
                continue
            if code:
                last_code = code
                q_code = _unique(code)
            else:
                base = last_code or section_id
                blank_counters[base] = blank_counters.get(base, 0) + 1
                q_code = _unique(f"{base}_{blank_counters[base]}")
            remarks_lower = remarks.lower()
            if remarks_lower in ("yes/no", "reconciliation"):
                qtype = "yes_no"
            elif "attach" in text.lower() or "upload" in text.lower():
                qtype = "file_upload"
            else:
                qtype = "text"
            is_label = (
                (not code) and (not remarks) and "?" not in text and len(text) < 45
                and (text.isupper() or text.endswith(":") or len(text.split()) <= 3)
            )
            questions.append({
                "code": q_code,
                "text": text,
                "type": qtype,
                "is_input": not is_label,
                "roll_forward": order in ROLL_FORWARD_SECTION_ORDERS,
            })
        sections.append({"section_id": section_id, "order": order, "name": title, "sheet_name": sheet.name, "questions": questions})
    return sections


async def seed_audit_template(db, logger=None):
    existing = await db.audit_templates.find_one({"template_id": TEMPLATE_ID}, {"_id": 0, "template_id": 1})
    if existing:
        return
    sections = parse_template()
    doc = {
        "template_id": TEMPLATE_ID,
        "name": "Standard Statutory Audit",
        "version": 1,
        "sections": sections,
        "created_at": datetime.now(timezone.utc).isoformat(),
    }
    await db.audit_templates.insert_one(doc)
    if logger:
        total_q = sum(len(s["questions"]) for s in sections)
        logger.info(f"Seeded audit template '{TEMPLATE_ID}' with {len(sections)} sections, {total_q} questions")


def compute_progress(template, responses):
    section_progress = []
    total_q = 0
    total_done = 0
    for sec in template["sections"]:
        answerable = [q for q in sec["questions"] if q["is_input"]]
        done = sum(1 for q in answerable if (responses.get(q["code"]) or {}).get("value"))
        pct = round(done / len(answerable) * 100) if answerable else 100
        section_progress.append({
            "section_id": sec["section_id"], "name": sec["name"], "order": sec["order"],
            "total": len(answerable), "done": done, "pct": pct,
        })
        total_q += len(answerable)
        total_done += done
    overall_pct = round(total_done / total_q * 100) if total_q else 0
    return section_progress, overall_pct, total_done, total_q


def roll_forward_responses(template, source_responses):
    """Copy answers for questions flagged roll_forward=True from a previous audit instance."""
    copied = {}
    for sec in template["sections"]:
        if sec["order"] not in ROLL_FORWARD_SECTION_ORDERS:
            continue
        for q in sec["questions"]:
            src = source_responses.get(q["code"])
            if src and src.get("value"):
                copied[q["code"]] = {
                    "value": src.get("value"),
                    "remarks": src.get("remarks"),
                    "attachment_path": src.get("attachment_path"),
                    "attachment_name": src.get("attachment_name"),
                    "drive_link": src.get("drive_link"),
                    "flagged": False,
                    "rolled_forward": True,
                }
    return copied


def generate_xlsx_export(template, audit_doc):
    responses = audit_doc.get("responses", {})
    wb = Workbook()
    wb.remove(wb.active)
    title_font = Font(bold=True, size=12, color="0A1128")
    header_font = Font(bold=True, size=10, color="FFFFFF")
    header_fill = PatternFill(start_color="0A1128", end_color="0A1128", fill_type="solid")

    for sec in sorted(template["sections"], key=lambda s: s["order"]):
        sheet_name = re.sub(r'[\\/*?:\[\]]', "", sec["sheet_name"])[:31]
        ws = wb.create_sheet(title=sheet_name or sec["name"][:31])
        ws.append([sec["name"], "", ""])
        ws["A1"].font = title_font
        ws.append(["Sl No", "Notes", "Remarks / Answer"])
        for cell in ws[2]:
            cell.font = header_font
            cell.fill = header_fill
        for q in sec["questions"]:
            resp = responses.get(q["code"], {})
            if not q["is_input"]:
                ws.append(["", q["text"], ""])
                continue
            answer = resp.get("value", "") or ""
            if resp.get("attachment_name"):
                answer = f"{answer} [Attached: {resp['attachment_name']}]".strip()
            if resp.get("remarks"):
                answer = f"{answer} — {resp['remarks']}".strip(" —")
            ws.append([q["code"], q["text"], answer])
        ws.column_dimensions["A"].width = 8
        ws.column_dimensions["B"].width = 70
        ws.column_dimensions["C"].width = 45

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def generate_pdf_export(template, audit_doc, firm_name="Nair & Nelliyatt Chartered Accountants"):
    from fpdf import FPDF

    def safe(text):
        text = str(text or "")
        replacements = {"\u2026": "...", "\u2019": "'", "\u2018": "'", "\u201c": '"', "\u201d": '"', "\u2013": "-", "\u2014": "-", "\u00a0": " "}
        for k, v in replacements.items():
            text = text.replace(k, v)
        return text.encode("latin-1", "replace").decode("latin-1")

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()

    pdf.set_font("Helvetica", "B", 16)
    pdf.set_text_color(10, 17, 40)
    pdf.cell(0, 10, firm_name, new_x="LMARGIN", new_y="NEXT", align="C")
    pdf.set_font("Helvetica", "", 10)
    pdf.set_text_color(120, 120, 120)
    pdf.cell(0, 6, "Digital Audit Workbook Report", new_x="LMARGIN", new_y="NEXT", align="C")
    pdf.ln(4)
    pdf.set_draw_color(212, 175, 55)
    pdf.line(10, pdf.get_y(), 200, pdf.get_y())
    pdf.ln(6)

    pdf.set_font("Helvetica", "B", 11)
    pdf.set_text_color(10, 17, 40)
    pdf.cell(0, 7, safe(f"Client: {audit_doc.get('client_name', '-')}"), new_x="LMARGIN", new_y="NEXT")
    pdf.cell(0, 7, safe(f"Period: {audit_doc.get('period', '-')}"), new_x="LMARGIN", new_y="NEXT")
    pdf.cell(0, 7, safe(f"Status: {audit_doc.get('status', '-')}"), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)

    responses = audit_doc.get("responses", {})
    for sec in sorted(template["sections"], key=lambda s: s["order"]):
        pdf.set_x(pdf.l_margin)
        pdf.set_font("Helvetica", "B", 11)
        pdf.set_text_color(212, 175, 55)
        pdf.multi_cell(0, 7, safe(sec["name"]))
        pdf.ln(1)
        for q in sec["questions"]:
            if not q["is_input"]:
                pdf.set_x(pdf.l_margin)
                pdf.set_font("Helvetica", "B", 9)
                pdf.set_text_color(80, 80, 80)
                pdf.multi_cell(0, 5, safe(f"  {q['text']}"))
                continue
            resp = responses.get(q["code"], {})
            answer = resp.get("value") or "-"
            pdf.set_x(pdf.l_margin)
            pdf.set_font("Helvetica", "", 9)
            pdf.set_text_color(30, 30, 30)
            pdf.multi_cell(0, 5, safe(f"  [{q['code']}] {q['text']}"))
            pdf.set_x(pdf.l_margin)
            pdf.set_font("Helvetica", "I", 9)
            flag = "  (FLAGGED)" if resp.get("flagged") else ""
            if resp.get("flagged"):
                pdf.set_text_color(150, 20, 20)
            else:
                pdf.set_text_color(20, 110, 60)
            attach_note = f" [Attached: {resp['attachment_name']}]" if resp.get("attachment_name") else ""
            pdf.multi_cell(0, 5, safe(f"      Answer: {answer}{attach_note}{flag}"))
        pdf.ln(2)

    pdf.set_font("Helvetica", "I", 8)
    pdf.set_text_color(150, 150, 150)
    pdf.cell(0, 5, f"Generated on {datetime.now(timezone.utc).strftime('%d %b %Y at %H:%M UTC')}", new_x="LMARGIN", new_y="NEXT", align="C")

    buffer = io.BytesIO()
    pdf.output(buffer)
    buffer.seek(0)
    return buffer
