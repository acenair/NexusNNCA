from fastapi import FastAPI, APIRouter, HTTPException, Header, Query, UploadFile, File, Response, Cookie
from fastapi.responses import StreamingResponse
from dotenv import load_dotenv
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
import os
import logging
from pathlib import Path
from pydantic import BaseModel, Field, ConfigDict
from typing import List, Optional, Dict, Any
import uuid
from datetime import datetime, timezone, timedelta
import bcrypt
import jwt
import requests
from emergentintegrations.llm.chat import LlmChat, UserMessage
import json
import asyncio
import io

# Import shared core
from core import (
    db, logger, EMERGENT_LLM_KEY, JWT_SECRET, JWT_ALGORITHM,
    VAPID_PRIVATE_KEY, VAPID_PUBLIC_KEY, VAPID_CLAIMS_EMAIL,
    STORAGE_URL, APP_NAME,
    init_storage, put_object, get_object,
    hash_password, verify_password, create_jwt_token,
    get_current_user, require_partner, log_activity,
)

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

# Create the main app (shared deps imported from core.py)
app = FastAPI()
api_router = APIRouter(prefix="/api")

# ============= MODELS =============

class User(BaseModel):
    model_config = ConfigDict(extra="ignore")
    user_id: str
    email: str
    name: str
    title: Optional[str] = None
    role: Optional[str] = None
    picture: Optional[str] = None
    created_at: str

class UserSession(BaseModel):
    model_config = ConfigDict(extra="ignore")
    user_id: str
    session_token: str
    expires_at: str
    created_at: str

class Client(BaseModel):
    model_config = ConfigDict(extra="ignore")
    client_id: str
    name: str
    entity_type: str
    jurisdiction: Optional[str] = None
    trade_licence_no: Optional[str] = None
    trn: Optional[str] = None
    ct_registration_no: Optional[str] = None
    vat_registration_date: Optional[str] = None
    tax_period: Optional[str] = None
    aml_risk_rating: Optional[str] = None
    pep_flag: bool = False
    relationship_manager_id: Optional[str] = None
    status: str = "Active"
    created_at: str
    active_services: List[str] = []

class Task(BaseModel):
    model_config = ConfigDict(extra="ignore")
    task_id: str
    title: str
    description: Optional[str] = None
    service_module: str
    client_id: Optional[str] = None
    due_date: str
    priority: str
    assigned_to: str
    status: str = "Pending"
    created_by: str
    created_at: str
    completed_at: Optional[str] = None

class Activity(BaseModel):
    model_config = ConfigDict(extra="ignore")
    activity_id: str
    event_type: str
    description: str
    client_id: Optional[str] = None
    client_name: Optional[str] = None
    user_id: str
    created_at: str

class ChatMessage(BaseModel):
    model_config = ConfigDict(extra="ignore")
    message_id: str
    session_id: str
    user_id: str
    role: str
    content: str
    created_at: str


# Storage/Auth helpers imported from core.py

# ============= AUTH ROUTES =============

@api_router.post("/auth/register")
async def register(email: str, password: str, name: str):
    existing = await db.users.find_one({"email": email})
    if existing:
        raise HTTPException(status_code=400, detail="Email already registered")
    
    user_id = f"user_{uuid.uuid4().hex[:12]}"
    hashed_pwd = hash_password(password)
    
    user_doc = {
        "user_id": user_id,
        "email": email,
        "name": name,
        "password": hashed_pwd,
        "role": "staff",
        "status": "pending_approval",
        "picture": None,
        "created_at": datetime.now(timezone.utc).isoformat()
    }
    await db.users.insert_one(user_doc)
    
    return {"message": "Account created. Pending admin approval.", "error": "pending_approval"}

class LoginRequest(BaseModel):
    email: str
    password: str

@api_router.post("/auth/login")
async def login(req: LoginRequest):
    user_doc = await db.users.find_one({"email": req.email})
    if not user_doc or not verify_password(req.password, user_doc.get("password", "")):
        raise HTTPException(status_code=401, detail="Invalid credentials")
    
    session_token = create_jwt_token(user_doc["user_id"])
    session_doc = {
        "user_id": user_doc["user_id"],
        "session_token": session_token,
        "expires_at": (datetime.now(timezone.utc) + timedelta(days=7)).isoformat(),
        "created_at": datetime.now(timezone.utc).isoformat()
    }
    await db.user_sessions.insert_one(session_doc)
    
    return {
        "user_id": user_doc["user_id"],
        "email": user_doc["email"],
        "name": user_doc["name"],
        "title": user_doc.get("title"),
        "role": user_doc.get("role", "staff"),
        "picture": user_doc.get("picture"),
        "session_token": session_token
    }

@api_router.get("/auth/users-list")
async def get_users_list():
    users = await db.users.find({"status": {"$ne": "pending_approval"}}, {"_id": 0, "password": 0}).to_list(100)
    return [{"name": u.get("name"), "email": u.get("email"), "role": u.get("role", "staff"), "title": u.get("title", "")} for u in users]

@api_router.post("/auth/session")
async def create_session(session_id: str):
    # REMINDER: DO NOT HARDCODE THE URL, OR ADD ANY FALLBACKS OR REDIRECT URLS, THIS BREAKS THE AUTH
    try:
        resp = requests.get(
            "https://demobackend.emergentagent.com/auth/v1/env/oauth/session-data",
            headers={"X-Session-ID": session_id},
            timeout=10
        )
        resp.raise_for_status()
        data = resp.json()
    except Exception as e:
        logger.error(f"Failed to get session data: {e}")
        raise HTTPException(status_code=400, detail="Invalid session ID")
    
    user_doc = await db.users.find_one({"email": data["email"]})
    if user_doc:
        user_id = user_doc["user_id"]
        await db.users.update_one(
            {"user_id": user_id},
            {"$set": {"name": data["name"], "picture": data["picture"]}}
        )
        # Check if user is approved
        if user_doc.get("status") == "pending_approval":
            return {"error": "pending_approval", "message": "Your account is pending admin approval. Please contact the administrator."}
    else:
        user_id = f"user_{uuid.uuid4().hex[:12]}"
        user_doc = {
            "user_id": user_id,
            "email": data["email"],
            "name": data["name"],
            "picture": data["picture"],
            "role": "staff",
            "title": "",
            "status": "pending_approval",
            "created_at": datetime.now(timezone.utc).isoformat()
        }
        await db.users.insert_one(user_doc)
        return {"error": "pending_approval", "message": "Account created! Your access is pending admin approval."}
    
    session_doc = {
        "user_id": user_id,
        "session_token": data["session_token"],
        "expires_at": (datetime.now(timezone.utc) + timedelta(days=7)).isoformat(),
        "created_at": datetime.now(timezone.utc).isoformat()
    }
    await db.user_sessions.insert_one(session_doc)
    
    return {
        "user_id": user_id,
        "email": data["email"],
        "name": data["name"],
        "picture": data["picture"],
        "session_token": data["session_token"]
    }

@api_router.get("/auth/me")
async def get_me(authorization: str = Header(None), session_token: str = Cookie(None)):
    user = await get_current_user(authorization, session_token)
    return user

@api_router.post("/auth/logout")
async def logout(authorization: str = Header(None), session_token: str = Cookie(None)):
    # Prefer Authorization header over cookie (consistent with get_current_user)
    token = None
    if authorization and authorization.startswith("Bearer "):
        token = authorization.split(" ")[1]
    elif session_token:
        token = session_token
    if token:
        await db.user_sessions.delete_one({"session_token": token})
    return {"message": "Logged out"}

# ============= DASHBOARD ROUTES =============

@api_router.get("/dashboard/stats")
async def get_dashboard_stats(authorization: str = Header(None), session_token: str = Cookie(None)):
    user = await get_current_user(authorization, session_token)
    
    active_clients = await db.clients.count_documents({"status": "Active"})
    total_tasks = await db.tasks.count_documents({"status": {"$ne": "Completed"}})
    overdue_tasks = await db.tasks.count_documents({
        "status": {"$ne": "Completed"},
        "due_date": {"$lt": datetime.now(timezone.utc).isoformat()}
    })
    
    current_month_start = datetime.now(timezone.utc).replace(day=1, hour=0, minute=0, second=0, microsecond=0).isoformat()
    filings_this_month = await db.vat_filings.count_documents({
        "filing_due_date": {"$gte": current_month_start}
    })
    
    aml_alerts = await db.aml_alerts.count_documents({"status": "Flagged"})
    
    return {
        "active_clients": active_clients,
        "open_tasks": total_tasks,
        "overdue_tasks": overdue_tasks,
        "filings_this_month": filings_this_month,
        "aml_alerts": aml_alerts
    }

@api_router.get("/dashboard/activities")
async def get_recent_activities(authorization: str = Header(None), session_token: str = Cookie(None), limit: int = 20):
    user = await get_current_user(authorization, session_token)
    activities = await db.activities.find({}, {"_id": 0}).sort("created_at", -1).limit(limit).to_list(limit)
    return activities

# ============= CLIENT ROUTES =============

@api_router.get("/clients")
async def get_clients(authorization: str = Header(None), session_token: str = Cookie(None)):
    user = await get_current_user(authorization, session_token)
    # Client role users can only see their own linked client record
    if user.get("role") == "client":
        client_id = user.get("client_id")
        if not client_id:
            return []
        client = await db.clients.find_one({"client_id": client_id}, {"_id": 0})
        return [client] if client else []
    clients = await db.clients.find({}, {"_id": 0}).sort("created_at", -1).to_list(1000)
    return clients

@api_router.post("/clients")
async def create_client(name: str, entity_type: str, authorization: str = Header(None), session_token: str = Cookie(None),
                       jurisdiction: Optional[str] = None, trade_licence_no: Optional[str] = None,
                       trn: Optional[str] = None, aml_risk_rating: Optional[str] = "Low"):
    user = await get_current_user(authorization, session_token)
    
    client_id = f"client_{uuid.uuid4().hex[:12]}"
    client_doc = {
        "client_id": client_id,
        "name": name,
        "entity_type": entity_type,
        "jurisdiction": jurisdiction,
        "trade_licence_no": trade_licence_no,
        "trn": trn,
        "aml_risk_rating": aml_risk_rating,
        "pep_flag": False,
        "status": "Active",
        "active_services": [],
        "created_at": datetime.now(timezone.utc).isoformat()
    }
    await db.clients.insert_one(client_doc)
    await log_activity("Client added", f"Added new client: {name}", user["user_id"], client_id, name)
    
    # Return without _id
    return {
        "client_id": client_id,
        "name": name,
        "entity_type": entity_type,
        "jurisdiction": jurisdiction,
        "trade_licence_no": trade_licence_no,
        "trn": trn,
        "aml_risk_rating": aml_risk_rating,
        "pep_flag": False,
        "status": "Active",
        "active_services": [],
        "created_at": client_doc["created_at"]
    }

@api_router.get("/clients/{client_id}")
async def get_client(client_id: str, authorization: str = Header(None), session_token: str = Cookie(None)):
    user = await get_current_user(authorization, session_token)
    client = await db.clients.find_one({"client_id": client_id}, {"_id": 0})
    if not client:
        raise HTTPException(status_code=404, detail="Client not found")
    return client

class UpdateClientRequest(BaseModel):
    name: Optional[str] = None
    entity_type: Optional[str] = None
    jurisdiction: Optional[str] = None
    trade_licence_no: Optional[str] = None
    trn: Optional[str] = None
    ct_registration_no: Optional[str] = None
    vat_registration_date: Optional[str] = None
    tax_period: Optional[str] = None
    aml_risk_rating: Optional[str] = None
    pep_flag: Optional[bool] = None
    status: Optional[str] = None
    active_services: Optional[List[str]] = None
    relationship_manager_id: Optional[str] = None

@api_router.patch("/clients/{client_id}")
async def update_client(client_id: str, req: UpdateClientRequest, authorization: str = Header(None), session_token: str = Cookie(None)):
    user = await get_current_user(authorization, session_token)
    if user.get("title") != "Managing Partner":
        raise HTTPException(status_code=403, detail="Only the Managing Partner can edit clients")
    update_data = {k: v for k, v in req.model_dump().items() if v is not None}
    if not update_data:
        raise HTTPException(status_code=400, detail="No fields to update")
    result = await db.clients.update_one({"client_id": client_id}, {"$set": update_data})
    if result.matched_count == 0:
        raise HTTPException(status_code=404, detail="Client not found")
    updated = await db.clients.find_one({"client_id": client_id}, {"_id": 0})
    await log_activity("Client updated", f"Updated client: {updated.get('name', client_id)}", user["user_id"], client_id, updated.get("name"))
    return updated

# ============= TASK ROUTES =============

@api_router.get("/tasks")
async def get_tasks(authorization: str = Header(None), session_token: str = Cookie(None), filter: Optional[str] = None):
    user = await get_current_user(authorization, session_token)
    
    query = {}
    if filter == "my":
        query["assigned_to"] = user["user_id"]
    elif filter == "overdue":
        query["status"] = {"$ne": "Completed"}
        query["due_date"] = {"$lt": datetime.now(timezone.utc).isoformat()}
    elif filter == "week":
        week_end = (datetime.now(timezone.utc) + timedelta(days=7)).isoformat()
        query["status"] = {"$ne": "Completed"}
        query["due_date"] = {"$lte": week_end}
    
    tasks = await db.tasks.find(query, {"_id": 0}).sort("due_date", 1).to_list(1000)
    return tasks

class CreateTaskRequest(BaseModel):
    title: str
    service_module: str
    due_date: str
    priority: str
    description: Optional[str] = None
    client_id: Optional[str] = None
    client_name: Optional[str] = None
    assigned_to: Optional[str] = None
    assigned_to_name: Optional[str] = None

@api_router.post("/tasks")
async def create_task(req: CreateTaskRequest,
                     authorization: str = Header(None), session_token: str = Cookie(None)):
    user = await get_current_user(authorization, session_token)
    
    task_id = f"task_{uuid.uuid4().hex[:12]}"
    task_doc = {
        "task_id": task_id,
        "title": req.title,
        "description": req.description,
        "service_module": req.service_module,
        "client_id": req.client_id,
        "client_name": req.client_name,
        "due_date": req.due_date,
        "priority": req.priority,
        "assigned_to": req.assigned_to or user["user_id"],
        "assigned_to_name": req.assigned_to_name,
        "status": "Pending",
        "created_by": user["user_id"],
        "created_by_name": user.get("name"),
        "created_at": datetime.now(timezone.utc).isoformat()
    }
    await db.tasks.insert_one(task_doc)
    await log_activity("Task created", f"Created task: {req.title}", user["user_id"], req.client_id, req.client_name)
    
    task_doc.pop("_id", None)
    return task_doc

@api_router.patch("/tasks/{task_id}")
async def update_task(task_id: str, status: Optional[str] = None, authorization: str = Header(None), session_token: str = Cookie(None)):
    user = await get_current_user(authorization, session_token)
    
    update_data = {}
    if status:
        update_data["status"] = status
        if status == "Completed":
            update_data["completed_at"] = datetime.now(timezone.utc).isoformat()
    
    await db.tasks.update_one({"task_id": task_id}, {"$set": update_data})
    return {"message": "Task updated"}

# ============= BULK TASK UPLOAD =============

TASK_HEADER_MAP = {
    # title mappings
    "title": "title", "task title": "title", "task name": "title", "task": "title",
    "name": "title", "subject": "title", "description_title": "title",
    # description mappings
    "description": "description", "desc": "description", "details": "description",
    "notes": "description", "note": "description", "remarks": "description",
    # service module mappings
    "service module": "service_module", "service_module": "service_module",
    "service": "service_module", "module": "service_module", "category": "service_module",
    "department": "service_module", "type": "service_module",
    # client mappings
    "client": "client_name", "client name": "client_name", "client_name": "client_name",
    "company": "client_name", "entity": "client_name", "customer": "client_name",
    # assigned to mappings
    "assigned to": "assigned_to_name", "assigned_to": "assigned_to_name",
    "assignee": "assigned_to_name", "assigned": "assigned_to_name",
    "owner": "assigned_to_name", "responsible": "assigned_to_name",
    "staff": "assigned_to_name", "team member": "assigned_to_name",
    # due date mappings
    "due date": "due_date", "due_date": "due_date", "deadline": "due_date",
    "due": "due_date", "target date": "due_date", "end date": "due_date",
    # priority mappings
    "priority": "priority", "urgency": "priority", "importance": "priority",
    # status mappings
    "status": "status", "state": "status", "progress": "status",
}

def parse_date_flexible(val):
    """Try to parse various date formats into YYYY-MM-DD."""
    if not val:
        return None
    val = str(val).strip()
    for fmt in ["%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y", "%d %b %Y", "%d %B %Y", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d %H:%M:%S"]:
        try:
            return datetime.strptime(val, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    # Try pandas-like parsing
    try:
        from dateutil import parser as dateparser
        return dateparser.parse(val).strftime("%Y-%m-%d")
    except Exception:
        pass
    return val

@api_router.post("/tasks/bulk-upload")
async def bulk_upload_tasks(file: UploadFile = File(...), authorization: str = Header(None), session_token: str = Cookie(None)):
    """Parse CSV or XLSX file, auto-map columns, and create tasks in bulk."""
    user = await get_current_user(authorization, session_token)
    if user.get("role") != "partner":
        raise HTTPException(status_code=403, detail="Only partners can bulk upload tasks")

    # Read file content
    content = await file.read()
    filename = file.filename or ""

    rows = []
    headers = []

    if filename.lower().endswith(".xlsx") or filename.lower().endswith(".xls"):
        import openpyxl
        import io
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        ws = wb.active
        all_rows = list(ws.iter_rows(values_only=True))
        if len(all_rows) < 2:
            raise HTTPException(status_code=400, detail="File has no data rows")
        headers = [str(h).strip() if h else "" for h in all_rows[0]]
        for row in all_rows[1:]:
            rows.append([str(cell) if cell is not None else "" for cell in row])
    elif filename.lower().endswith(".csv"):
        import csv
        import io
        text = content.decode("utf-8-sig")
        reader = csv.reader(io.StringIO(text))
        all_rows = list(reader)
        if len(all_rows) < 2:
            raise HTTPException(status_code=400, detail="File has no data rows")
        headers = [h.strip() for h in all_rows[0]]
        rows = all_rows[1:]
    else:
        raise HTTPException(status_code=400, detail="Unsupported file format. Please upload CSV or XLSX.")

    # Auto-map headers
    column_map = {}
    for idx, h in enumerate(headers):
        normalized = h.lower().strip()
        if normalized in TASK_HEADER_MAP:
            column_map[idx] = TASK_HEADER_MAP[normalized]

    if "title" not in column_map.values():
        raise HTTPException(status_code=400, detail=f"Could not find a 'Title' column. Found headers: {headers}")

    # Resolve staff names to emails
    all_staff = await db.users.find({}, {"_id": 0, "name": 1, "email": 1}).to_list(100)
    staff_name_to_email = {u["name"].lower(): u["email"] for u in all_staff}

    # Resolve client names to IDs
    all_clients = await db.clients.find({}, {"_id": 0, "client_id": 1, "name": 1}).to_list(500)
    client_name_to_id = {c["name"].lower(): c["client_id"] for c in all_clients}

    created = []
    errors = []
    for row_idx, row in enumerate(rows, start=2):
        try:
            task_data = {}
            for col_idx, field in column_map.items():
                if col_idx < len(row):
                    task_data[field] = row[col_idx].strip()

            title = task_data.get("title", "").strip()
            if not title:
                continue  # Skip empty rows

            # Parse due date
            due_date = parse_date_flexible(task_data.get("due_date"))
            if not due_date:
                due_date = (datetime.now(timezone.utc) + timedelta(days=7)).strftime("%Y-%m-%d")

            # Resolve staff
            assignee_name = task_data.get("assigned_to_name", "").strip()
            assignee_email = staff_name_to_email.get(assignee_name.lower(), "")

            # Resolve client
            client_name = task_data.get("client_name", "").strip()
            client_id = client_name_to_id.get(client_name.lower(), "")

            # Normalize priority
            raw_priority = task_data.get("priority", "Medium").strip().capitalize()
            priority = raw_priority if raw_priority in ("High", "Medium", "Low") else "Medium"

            # Normalize status
            raw_status = task_data.get("status", "Pending").strip()
            status = raw_status if raw_status in ("Pending", "In Progress", "Completed") else "Pending"

            service_module = task_data.get("service_module", "General").strip() or "General"
            description = task_data.get("description", "").strip()

            task_id = f"task_{uuid.uuid4().hex[:12]}"
            task_doc = {
                "task_id": task_id,
                "title": title,
                "description": description or f"Bulk imported from {filename}",
                "service_module": service_module,
                "client_id": client_id,
                "client_name": client_name,
                "due_date": due_date,
                "priority": priority,
                "assigned_to": assignee_email or user.get("email"),
                "assigned_to_name": assignee_name or user.get("name"),
                "status": status,
                "created_by": user["user_id"],
                "created_by_name": user.get("name"),
                "created_at": datetime.now(timezone.utc).isoformat(),
                "source": "bulk_upload",
            }
            await db.tasks.insert_one(task_doc)
            created.append({"row": row_idx, "task_id": task_id, "title": title})
        except Exception as e:
            errors.append({"row": row_idx, "error": str(e)})

    await log_activity("Bulk task upload", f"Uploaded {len(created)} tasks from {filename}", user["user_id"])

    return {
        "total_rows": len(rows),
        "created": len(created),
        "errors": len(errors),
        "error_details": errors[:20],
        "mapped_columns": {headers[idx]: field for idx, field in column_map.items()},
        "message": f"Successfully imported {len(created)} task(s) from {filename}",
    }

@api_router.post("/tasks/bulk-preview")
async def preview_bulk_upload(file: UploadFile = File(...), authorization: str = Header(None), session_token: str = Cookie(None)):
    """Preview column mapping and first few rows without importing."""
    user = await get_current_user(authorization, session_token)
    if user.get("role") != "partner":
        raise HTTPException(status_code=403, detail="Only partners can bulk upload tasks")

    content = await file.read()
    filename = file.filename or ""

    rows = []
    headers = []

    if filename.lower().endswith(".xlsx") or filename.lower().endswith(".xls"):
        import openpyxl
        import io
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        ws = wb.active
        all_rows = list(ws.iter_rows(values_only=True))
        if len(all_rows) < 1:
            raise HTTPException(status_code=400, detail="File is empty")
        headers = [str(h).strip() if h else "" for h in all_rows[0]]
        for row in all_rows[1:6]:
            rows.append([str(cell) if cell is not None else "" for cell in row])
    elif filename.lower().endswith(".csv"):
        import csv
        import io
        text = content.decode("utf-8-sig")
        reader = csv.reader(io.StringIO(text))
        all_rows = list(reader)
        if len(all_rows) < 1:
            raise HTTPException(status_code=400, detail="File is empty")
        headers = [h.strip() for h in all_rows[0]]
        rows = all_rows[1:6]
    else:
        raise HTTPException(status_code=400, detail="Unsupported file format. Please upload CSV or XLSX.")

    column_map = {}
    unmapped = []
    for idx, h in enumerate(headers):
        normalized = h.lower().strip()
        if normalized in TASK_HEADER_MAP:
            column_map[h] = TASK_HEADER_MAP[normalized]
        else:
            unmapped.append(h)

    return {
        "filename": filename,
        "total_rows": len(all_rows) - 1 if filename.lower().endswith(".csv") else len(list(rows)),
        "headers": headers,
        "mapped_columns": column_map,
        "unmapped_columns": unmapped,
        "preview_rows": rows,
    }


# ============= EVENTS / CALENDAR ROUTES =============

class CreateEventRequest(BaseModel):
    title: str
    event_type: str  # meeting, followup, task, deadline
    date: str
    time: Optional[str] = None
    client_name: Optional[str] = None
    client_id: Optional[str] = None
    notes: Optional[str] = None
    assigned_to: Optional[str] = None
    assigned_to_name: Optional[str] = None

@api_router.post("/events")
async def create_event(req: CreateEventRequest, authorization: str = Header(None), session_token: str = Cookie(None)):
    user = await get_current_user(authorization, session_token)
    event_id = f"evt_{uuid.uuid4().hex[:12]}"
    event_doc = {
        "event_id": event_id,
        "title": req.title,
        "event_type": req.event_type,
        "date": req.date,
        "time": req.time,
        "client_name": req.client_name,
        "client_id": req.client_id,
        "notes": req.notes,
        "assigned_to": req.assigned_to,
        "assigned_to_name": req.assigned_to_name,
        "created_by": user["user_id"],
        "created_by_name": user.get("name"),
        "created_at": datetime.now(timezone.utc).isoformat()
    }
    await db.events.insert_one(event_doc)
    label = req.event_type.replace("followup", "follow-up").capitalize()
    await log_activity(f"{label} created", f"{label}: {req.title}", user["user_id"], req.client_id, req.client_name)
    event_doc.pop("_id", None)
    return event_doc

@api_router.get("/events")
async def get_events(authorization: str = Header(None), session_token: str = Cookie(None), month: Optional[str] = None):
    user = await get_current_user(authorization, session_token)
    query = {}
    if month:
        query["date"] = {"$regex": f"^{month}"}
    events = await db.events.find(query, {"_id": 0}).sort("date", 1).to_list(500)
    return events

@api_router.delete("/events/{event_id}")
async def delete_event(event_id: str, authorization: str = Header(None), session_token: str = Cookie(None)):
    user = await get_current_user(authorization, session_token)
    await db.events.delete_one({"event_id": event_id})
    return {"message": "Event deleted"}

# ============= APPRECIATION ROUTES =============

class CreateAppreciationRequest(BaseModel):
    staff_email: str
    staff_name: str
    categories: List[str]
    rating: int
    month: str
    message: Optional[str] = None

@api_router.post("/appreciations")
async def create_appreciation(req: CreateAppreciationRequest, authorization: str = Header(None), session_token: str = Cookie(None)):
    user = await get_current_user(authorization, session_token)
    if user.get("role") != "partner":
        raise HTTPException(status_code=403, detail="Only partners can give appreciations")
    appr_id = f"appr_{uuid.uuid4().hex[:12]}"
    appr_doc = {
        "appreciation_id": appr_id,
        "staff_email": req.staff_email,
        "staff_name": req.staff_name,
        "categories": req.categories,
        "rating": req.rating,
        "month": req.month,
        "message": req.message,
        "given_by": user["user_id"],
        "given_by_name": user.get("name"),
        "created_at": datetime.now(timezone.utc).isoformat()
    }
    await db.appreciations.insert_one(appr_doc)
    await log_activity("Appreciation given", f"Appreciated {req.staff_name}: {', '.join(req.categories)}", user["user_id"])
    appr_doc.pop("_id", None)
    return appr_doc

@api_router.get("/appreciations")
async def get_appreciations(authorization: str = Header(None), session_token: str = Cookie(None), staff_email: Optional[str] = None):
    user = await get_current_user(authorization, session_token)
    query = {}
    if staff_email:
        query["staff_email"] = staff_email
    apprs = await db.appreciations.find(query, {"_id": 0}).sort("created_at", -1).to_list(200)
    return apprs


# ============= AI ASSISTANT ROUTES =============

class AIChatRequest(BaseModel):
    session_id: str
    message: str

async def get_firm_context(user):
    """Gather firm data context for AI to reference."""
    clients = await db.clients.find({}, {"_id": 0, "name": 1, "status": 1, "active_services": 1, "jurisdiction": 1}).to_list(50)
    tasks = await db.tasks.find({}, {"_id": 0, "title": 1, "status": 1, "assigned_to_name": 1, "client_name": 1, "due_date": 1, "service_module": 1, "priority": 1}).sort("due_date", 1).to_list(50)
    events = await db.events.find({}, {"_id": 0, "title": 1, "event_type": 1, "date": 1, "client_name": 1}).sort("date", 1).to_list(30)
    staff = await db.users.find({"role": "staff"}, {"_id": 0, "name": 1, "title": 1, "email": 1}).to_list(20)

    lines = ["=== FIRM DATA (Nair & Nelliyatt Chartered Accountants) ==="]
    if clients:
        lines.append(f"\nClients ({len(clients)}):")
        for c in clients:
            svcs = ', '.join(c.get('active_services', []))
            lines.append(f"  - {c['name']} | {c.get('status','N/A')} | {c.get('jurisdiction','')} | Services: {svcs}")
    if tasks:
        lines.append(f"\nActive Tasks ({len(tasks)}):")
        for t in tasks:
            lines.append(f"  - {t.get('title','')} | Client: {t.get('client_name','N/A')} | Assigned: {t.get('assigned_to_name','N/A')} | Due: {t.get('due_date','N/A')} | Status: {t.get('status','Pending')} | Priority: {t.get('priority','Medium')}")
    if events:
        lines.append(f"\nUpcoming Events ({len(events)}):")
        for e in events:
            lines.append(f"  - {e.get('date','')} | {e.get('event_type','')} | {e.get('title','')} | Client: {e.get('client_name','')}")
    if staff:
        lines.append(f"\nStaff ({len(staff)}):")
        for s in staff:
            lines.append(f"  - {s['name']} ({s.get('title','Staff')})")
    return '\n'.join(lines)

@api_router.post("/ai/chat")
async def chat_with_ai(req: AIChatRequest, authorization: str = Header(None), session_token: str = Cookie(None)):
    user = await get_current_user(authorization, session_token)

    # Save user message
    user_msg_doc = {
        "message_id": f"msg_{uuid.uuid4().hex[:12]}",
        "session_id": req.session_id,
        "user_id": user["user_id"],
        "role": "user",
        "content": req.message,
        "created_at": datetime.now(timezone.utc).isoformat()
    }
    await db.chat_messages.insert_one(user_msg_doc)

    # Build context from firm data
    firm_context = await get_firm_context(user)

    # Build chat history for context
    history = await db.chat_messages.find(
        {"session_id": req.session_id},
        {"_id": 0, "role": 1, "content": 1}
    ).sort("created_at", 1).to_list(40)
    history_text = ""
    if len(history) > 1:
        recent = history[-10:]
        history_text = "\n\n=== RECENT CONVERSATION ===\n" + "\n".join([f"{m['role'].upper()}: {m['content']}" for m in recent[:-1]])

    system_message = f"""You are the AI Compliance Assistant for Nair & Nelliyatt Chartered Accountants, a UAE-based audit and advisory firm.

You serve two roles:
1. UAE REGULATORY EXPERT — Answer questions about:
   - UAE VAT (Federal Decree-Law No. 8/2017, amended)
   - UAE Corporate Tax (Federal Decree-Law No. 47/2022, 9% rate, AED 375K threshold)
   - AML/CFT (Federal Law No. 20/2018 — DNFBP obligations, goAML, STR filing)
   - UAE Companies Law (Federal Decree-Law No. 32/2021)
   - International audit standards (ISA), IFRS framework
   - FTA portal processes, TRN registration, voluntary disclosures

2. FIRM DATA ASSISTANT — You have real-time access to the firm's client portfolio, tasks, events, and staff. Answer questions about client status, deadlines, workload, audit progress, etc.

{firm_context}
{history_text}

Guidelines:
- Be concise and precise (under 300 words unless detail is requested)
- Cite applicable UAE law or standard when relevant
- For firm data questions, reference specific client names, dates, and staff assignments
- If you don't know something, say so clearly
- Format responses with markdown for clarity"""

    chat = LlmChat(
        api_key=EMERGENT_LLM_KEY,
        session_id=req.session_id,
        system_message=system_message
    ).with_model("gemini", "gemini-3-flash-preview")

    user_message = UserMessage(text=req.message)

    try:
        response = await chat.send_message(user_message)

        ai_msg_doc = {
            "message_id": f"msg_{uuid.uuid4().hex[:12]}",
            "session_id": req.session_id,
            "user_id": user["user_id"],
            "role": "assistant",
            "content": response,
            "created_at": datetime.now(timezone.utc).isoformat()
        }
        await db.chat_messages.insert_one(ai_msg_doc)

        return {"response": response, "session_id": req.session_id}
    except Exception as e:
        logger.error(f"AI chat error: {e}")
        raise HTTPException(status_code=500, detail=f"AI service error: {str(e)}")

@api_router.get("/ai/chat/{session_id}")
async def get_chat_history(session_id: str, authorization: str = Header(None), session_token: str = Cookie(None)):
    user = await get_current_user(authorization, session_token)
    messages = await db.chat_messages.find(
        {"session_id": session_id, "user_id": user["user_id"]},
        {"_id": 0}
    ).sort("created_at", 1).to_list(1000)
    return messages

@api_router.get("/ai/sessions")
async def get_chat_sessions(authorization: str = Header(None), session_token: str = Cookie(None)):
    user = await get_current_user(authorization, session_token)
    pipeline = [
        {"$match": {"user_id": user["user_id"]}},
        {"$group": {"_id": "$session_id", "last_msg": {"$last": "$content"}, "last_at": {"$last": "$created_at"}, "count": {"$sum": 1}}},
        {"$sort": {"last_at": -1}},
        {"$limit": 20}
    ]
    sessions = await db.chat_messages.aggregate(pipeline).to_list(20)
    return [{"session_id": s["_id"], "preview": s["last_msg"][:80] if s["last_msg"] else "", "last_at": s["last_at"], "message_count": s["count"]} for s in sessions]

@api_router.delete("/ai/chat/{session_id}")
async def delete_chat_session(session_id: str, authorization: str = Header(None), session_token: str = Cookie(None)):
    user = await get_current_user(authorization, session_token)
    await db.chat_messages.delete_many({"session_id": session_id, "user_id": user["user_id"]})
    return {"message": "Session deleted"}

# ============= FILE UPLOAD ROUTES =============

@api_router.post("/files/upload")
async def upload_file(file: UploadFile = File(...), authorization: str = Header(None), session_token: str = Cookie(None),
                     client_id: Optional[str] = None, document_type: Optional[str] = None):
    user = await get_current_user(authorization, session_token)
    
    ext = file.filename.split(".")[-1] if "." in file.filename else "bin"
    file_id = f"file_{uuid.uuid4().hex[:12]}"
    path = f"{APP_NAME}/uploads/{user['user_id']}/{file_id}.{ext}"
    
    data = await file.read()
    result = put_object(path, data, file.content_type or "application/octet-stream")
    
    file_doc = {
        "file_id": file_id,
        "storage_path": result["path"],
        "original_filename": file.filename,
        "content_type": file.content_type,
        "size": result["size"],
        "client_id": client_id,
        "document_type": document_type,
        "uploaded_by": user["user_id"],
        "is_deleted": False,
        "created_at": datetime.now(timezone.utc).isoformat()
    }
    await db.documents.insert_one(file_doc)
    
    return {"file_id": file_id, "filename": file.filename, "path": result["path"]}

@api_router.get("/documents")
async def list_documents(authorization: str = Header(None), session_token: str = Cookie(None),
                        client_id: Optional[str] = None, document_type: Optional[str] = None):
    user = await get_current_user(authorization, session_token)
    query = {"is_deleted": False}
    if client_id:
        query["client_id"] = client_id
    if document_type:
        query["document_type"] = document_type
    docs = await db.documents.find(query, {"_id": 0}).sort("created_at", -1).to_list(500)
    # Enrich with uploader name
    user_cache = {}
    for doc in docs:
        uid = doc.get("uploaded_by")
        if uid and uid not in user_cache:
            u = await db.users.find_one({"user_id": uid}, {"_id": 0, "name": 1})
            user_cache[uid] = u.get("name", "Unknown") if u else "Unknown"
        doc["uploaded_by_name"] = user_cache.get(uid, "Unknown")
    return docs

@api_router.delete("/documents/{file_id}")
async def delete_document(file_id: str, authorization: str = Header(None), session_token: str = Cookie(None)):
    user = await get_current_user(authorization, session_token)
    result = await db.documents.update_one({"file_id": file_id, "is_deleted": False}, {"$set": {"is_deleted": True, "deleted_at": datetime.now(timezone.utc).isoformat(), "deleted_by": user["user_id"]}})
    if result.matched_count == 0:
        raise HTTPException(status_code=404, detail="Document not found")
    return {"message": "Document deleted"}

@api_router.get("/files/{file_id}")
async def download_file(file_id: str, authorization: str = Header(None), session_token: str = Cookie(None), auth: str = Query(None)):
    auth_header = authorization or (f"Bearer {auth}" if auth else None)
    user = await get_current_user(auth_header, session_token)
    
    file_doc = await db.documents.find_one({"file_id": file_id, "is_deleted": {"$ne": True}}, {"_id": 0})
    if not file_doc:
        raise HTTPException(status_code=404, detail="File not found")
    
    # Support both object storage and direct MongoDB storage
    if file_doc.get("storage_path"):
        data, content_type = get_object(file_doc["storage_path"])
        return Response(content=data, media_type=file_doc.get("content_type", content_type))
    elif file_doc.get("data"):
        return Response(content=file_doc["data"], media_type=file_doc.get("content_type", "application/octet-stream"))
    else:
        raise HTTPException(status_code=404, detail="File data not found")

# Simplified VAT, Audit, Corporate, AML, Advisory routes (basic CRUD for MVP)

@api_router.get("/vat/registrations")
async def get_vat_registrations(authorization: str = Header(None), session_token: str = Cookie(None)):
    user = await get_current_user(authorization, session_token)
    registrations = await db.vat_registrations.find({}, {"_id": 0}).to_list(1000)
    return registrations

@api_router.post("/vat/registrations")
async def create_vat_registration(client_id: str, annual_turnover: float, registration_type: str,
                                 authorization: str = Header(None), session_token: str = Cookie(None)):
    user = await get_current_user(authorization, session_token)
    
    client = await db.clients.find_one({"client_id": client_id}, {"_id": 0})
    if not client:
        raise HTTPException(status_code=404, detail="Client not found")
    
    reg_id = f"vat_reg_{uuid.uuid4().hex[:12]}"
    reg_doc = {
        "registration_id": reg_id,
        "client_id": client_id,
        "client_name": client["name"],
        "annual_turnover": annual_turnover,
        "registration_type": registration_type,
        "status": "Pending Documents",
        "created_at": datetime.now(timezone.utc).isoformat()
    }
    await db.vat_registrations.insert_one(reg_doc)
    await log_activity("VAT registration", f"VAT registration started for {client['name']}", user["user_id"], client_id, client["name"])
    
    # Return without _id
    return {
        "registration_id": reg_id,
        "client_id": client_id,
        "client_name": client["name"],
        "annual_turnover": annual_turnover,
        "registration_type": registration_type,
        "status": "Pending Documents",
        "created_at": reg_doc["created_at"]
    }

@api_router.get("/vat/filings")
async def get_vat_filings(authorization: str = Header(None), session_token: str = Cookie(None)):
    user = await get_current_user(authorization, session_token)
    filings = await db.vat_filings.find({}, {"_id": 0}).to_list(1000)
    return filings

@api_router.post("/vat/filings")
async def create_vat_filing(client_id: str, tax_period: str, period_end_date: str, filing_due_date: str,
                           authorization: str = Header(None), session_token: str = Cookie(None)):
    user = await get_current_user(authorization, session_token)
    
    client = await db.clients.find_one({"client_id": client_id}, {"_id": 0})
    if not client:
        raise HTTPException(status_code=404, detail="Client not found")
    
    filing_id = f"vat_filing_{uuid.uuid4().hex[:12]}"
    filing_doc = {
        "filing_id": filing_id,
        "client_id": client_id,
        "client_name": client["name"],
        "tax_period": tax_period,
        "period_end_date": period_end_date,
        "filing_due_date": filing_due_date,
        "status": "Pending",
        "box_1a": 0,
        "box_1b": 0,
        "box_2": 0,
        "box_3": 0,
        "box_4": 0,
        "created_at": datetime.now(timezone.utc).isoformat()
    }
    await db.vat_filings.insert_one(filing_doc)
    await log_activity("VAT filing", f"VAT filing created for {client['name']}", user["user_id"], client_id, client["name"])
    
    # Return without _id
    return {
        "filing_id": filing_id,
        "client_id": client_id,
        "client_name": client["name"],
        "tax_period": tax_period,
        "period_end_date": period_end_date,
        "filing_due_date": filing_due_date,
        "status": "Pending",
        "box_1a": 0,
        "box_1b": 0,
        "box_2": 0,
        "box_3": 0,
        "box_4": 0,
        "created_at": filing_doc["created_at"]
    }

@api_router.get("/audit/engagements")
async def get_audit_engagements(authorization: str = Header(None), session_token: str = Cookie(None), engagement_type: Optional[str] = None):
    user = await get_current_user(authorization, session_token)
    query = {"engagement_type": engagement_type} if engagement_type else {}
    engagements = await db.audit_engagements.find(query, {"_id": 0}).to_list(1000)
    return engagements

@api_router.post("/audit/engagements")
async def create_audit_engagement(client_id: str, engagement_type: str, period_start: str, period_end: str,
                                 lead_auditor: str, authorization: str = Header(None), session_token: str = Cookie(None)):
    user = await get_current_user(authorization, session_token)
    
    client = await db.clients.find_one({"client_id": client_id}, {"_id": 0})
    if not client:
        raise HTTPException(status_code=404, detail="Client not found")
    
    engagement_id = f"audit_{uuid.uuid4().hex[:12]}"
    engagement_doc = {
        "engagement_id": engagement_id,
        "client_id": client_id,
        "client_name": client["name"],
        "engagement_type": engagement_type,
        "period_start": period_start,
        "period_end": period_end,
        "lead_auditor": lead_auditor,
        "phase": "Planning",
        "risk_level": "Medium",
        "status": "Active",
        "created_at": datetime.now(timezone.utc).isoformat()
    }
    await db.audit_engagements.insert_one(engagement_doc)
    await log_activity("Audit engagement", f"{engagement_type} audit started for {client['name']}", user["user_id"], client_id, client["name"])
    
    # Return without _id
    return {
        "engagement_id": engagement_id,
        "client_id": client_id,
        "client_name": client["name"],
        "engagement_type": engagement_type,
        "period_start": period_start,
        "period_end": period_end,
        "lead_auditor": lead_auditor,
        "phase": "Planning",
        "risk_level": "Medium",
        "status": "Active",
        "created_at": engagement_doc["created_at"]
    }

@api_router.get("/aml/alerts")
async def get_aml_alerts(authorization: str = Header(None), session_token: str = Cookie(None)):
    user = await get_current_user(authorization, session_token)
    alerts = await db.aml_alerts.find({}, {"_id": 0}).sort("created_at", -1).to_list(1000)
    return alerts

@api_router.post("/aml/alerts")
async def create_aml_alert(client_id: str, amount: float, flag_reason: str,
                          authorization: str = Header(None), session_token: str = Cookie(None)):
    user = await get_current_user(authorization, session_token)
    
    client = await db.clients.find_one({"client_id": client_id}, {"_id": 0})
    if not client:
        raise HTTPException(status_code=404, detail="Client not found")
    
    risk_score = 50
    if amount >= 100000:
        risk_score = 85
    elif amount >= 55000:
        risk_score = 65
    
    alert_id = f"aml_{uuid.uuid4().hex[:12]}"
    alert_doc = {
        "alert_id": alert_id,
        "client_id": client_id,
        "client_name": client["name"],
        "transaction_id": f"txn_{uuid.uuid4().hex[:8]}",
        "amount": amount,
        "flag_reason": flag_reason,
        "risk_score": risk_score,
        "assigned_to": user["user_id"],
        "status": "Flagged",
        "created_at": datetime.now(timezone.utc).isoformat()
    }
    await db.aml_alerts.insert_one(alert_doc)
    await log_activity("AML flag", f"Transaction flagged for {client['name']}: AED {amount:,.2f}", user["user_id"], client_id, client["name"])
    
    # Return without _id
    return {
        "alert_id": alert_id,
        "client_id": client_id,
        "client_name": client["name"],
        "transaction_id": alert_doc["transaction_id"],
        "amount": amount,
        "flag_reason": flag_reason,
        "risk_score": risk_score,
        "assigned_to": user["user_id"],
        "status": "Flagged",
        "created_at": alert_doc["created_at"]
    }

@api_router.patch("/aml/alerts/{alert_id}")
async def update_aml_alert(alert_id: str, status: str, notes: Optional[str] = None,
                          authorization: str = Header(None), session_token: str = Cookie(None)):
    user = await get_current_user(authorization, session_token)
    
    update_data = {"status": status}
    if notes:
        update_data["mlro_notes"] = notes
    update_data["actioned_at"] = datetime.now(timezone.utc).isoformat()
    update_data["actioned_by"] = user["user_id"]
    
    await db.aml_alerts.update_one({"alert_id": alert_id}, {"$set": update_data})
    return {"message": "AML alert updated"}

@api_router.get("/analytics/stats")
async def get_analytics_stats(authorization: str = Header(None), session_token: str = Cookie(None)):
    user = await get_current_user(authorization, session_token)
    
    total_clients = await db.clients.count_documents({})
    total_tasks = await db.tasks.count_documents({})
    completed_tasks = await db.tasks.count_documents({"status": "Completed"})
    
    vat_count = await db.vat_filings.count_documents({})
    audit_count = await db.audit_engagements.count_documents({})
    aml_count = await db.aml_alerts.count_documents({})
    
    return {
        "total_clients": total_clients,
        "total_tasks": total_tasks,
        "completed_tasks": completed_tasks,
        "completion_rate": round((completed_tasks / total_tasks * 100) if total_tasks > 0 else 0, 1),
        "vat_filings": vat_count,
        "audit_engagements": audit_count,
        "aml_alerts": aml_count
    }

# ============= SERVICE ENGAGEMENTS (Unified) =============

CHECKLIST_TEMPLATES = {
    "statutory_audit": [
        {"group": "Planning", "items": ["Engagement letter", "Risk assessment", "Audit plan", "Materiality memo", "Team briefing"]},
        {"group": "Fieldwork", "items": ["Revenue testing", "Expense sampling", "Bank confirmations", "Inventory count", "Related party review"]},
        {"group": "Review & Reporting", "items": ["Draft report", "Management letter", "Partner review", "Client sign-off", "Filing"]},
    ],
    "internal_audit": [
        {"group": "Scoping", "items": ["Scope definition", "Risk universe update", "Audit program"]},
        {"group": "Testing", "items": ["Control testing", "Walkthrough", "Sample testing", "Exception analysis"]},
        {"group": "Reporting", "items": ["Draft findings", "Management response", "Final report"]},
    ],
    "vat_filing": [
        {"group": "Preparation", "items": ["Data collection", "Purchase invoices review", "Sales invoices review", "Reconciliation"]},
        {"group": "Filing", "items": ["Return preparation", "Box amounts calculation", "Review & approval", "FTA portal submission"]},
        {"group": "Completion", "items": ["Payment confirmation", "Filing receipt archive", "Client notification"]},
    ],
    "vat_registration": [
        {"group": "Documentation", "items": ["Trade licence copy", "Passport / Emirates ID", "Bank letter", "Turnover evidence"]},
        {"group": "Submission", "items": ["FTA portal application", "Supporting docs upload", "Application review"]},
        {"group": "Completion", "items": ["TRN issued", "Certificate archived", "Client notified"]},
    ],
    "corporate_tax": [
        {"group": "Preparation", "items": ["Financial data collection", "Revenue classification", "Exempt income review", "Deduction analysis"]},
        {"group": "Computation", "items": ["Taxable income calculation", "Tax liability computation", "Small business relief check"]},
        {"group": "Filing", "items": ["CT return preparation", "Review & approval", "Portal submission", "Payment processing"]},
    ],
    "aml_review": [
        {"group": "Client Due Diligence", "items": ["KYC documentation", "Beneficial ownership check", "PEP screening", "Sanctions screening"]},
        {"group": "Transaction Monitoring", "items": ["Unusual transaction review", "Threshold analysis", "STR assessment"]},
        {"group": "Reporting", "items": ["goAML report preparation", "MLRO review", "Filing submission"]},
    ],
}

class CreateEngagementRequest(BaseModel):
    service_type: str
    client_id: str
    assigned_to: Optional[str] = None
    assigned_to_name: Optional[str] = None
    phase: Optional[str] = None
    notes: Optional[str] = None
    workflow_id: Optional[str] = None

@api_router.post("/service/engagements")
async def create_service_engagement(req: CreateEngagementRequest, authorization: str = Header(None), session_token: str = Cookie(None)):
    user = await get_current_user(authorization, session_token)
    client = await db.clients.find_one({"client_id": req.client_id}, {"_id": 0, "name": 1})
    if not client:
        raise HTTPException(status_code=404, detail="Client not found")

    # If workflow_id provided, use workflow steps as checklist
    checklist = []
    workflow_name = None
    if req.workflow_id:
        wf = await db.workflows.find_one({"workflow_id": req.workflow_id, "is_deleted": {"$ne": True}}, {"_id": 0})
        if wf:
            workflow_name = wf.get("name")
            # Convert workflow steps into checklist groups
            checklist = [{"group": step["name"], "items": [{"label": step.get("description") or step["name"], "done": False}]} for step in wf.get("steps", [])]

    # Fallback to CHECKLIST_TEMPLATES if no workflow provided or found
    if not checklist:
        template = CHECKLIST_TEMPLATES.get(req.service_type, [])
        checklist = [{"group": g["group"], "items": [{"label": item, "done": False} for item in g["items"]]} for g in template]

    eng_id = f"eng_{uuid.uuid4().hex[:12]}"
    eng_doc = {
        "engagement_id": eng_id,
        "service_type": req.service_type,
        "client_id": req.client_id,
        "client_name": client["name"],
        "assigned_to": req.assigned_to,
        "assigned_to_name": req.assigned_to_name,
        "status": "Active",
        "phase": req.phase or (checklist[0]["group"] if checklist else ""),
        "notes": req.notes,
        "checklist": checklist,
        "workflow_id": req.workflow_id,
        "workflow_name": workflow_name,
        "created_by": user["user_id"],
        "created_by_name": user.get("name"),
        "created_at": datetime.now(timezone.utc).isoformat(),
    }
    await db.service_engagements.insert_one(eng_doc)
    eng_doc.pop("_id", None)
    await log_activity("Engagement created", f"{req.service_type.replace('_',' ').title()} for {client['name']}", user["user_id"], req.client_id, client["name"])
    return eng_doc

@api_router.get("/service/engagements")
async def list_service_engagements(authorization: str = Header(None), session_token: str = Cookie(None), service_type: Optional[str] = None, client_id: Optional[str] = None):
    user = await get_current_user(authorization, session_token)
    query = {"status": {"$ne": "Deleted"}}
    if service_type:
        query["service_type"] = service_type
    if client_id:
        query["client_id"] = client_id
    engs = await db.service_engagements.find(query, {"_id": 0}).sort("created_at", -1).to_list(200)
    return engs

class ToggleChecklistRequest(BaseModel):
    group_index: int
    item_index: int
    done: bool

@api_router.patch("/service/engagements/{engagement_id}/checklist")
async def toggle_checklist_item(engagement_id: str, req: ToggleChecklistRequest, authorization: str = Header(None), session_token: str = Cookie(None)):
    user = await get_current_user(authorization, session_token)
    eng = await db.service_engagements.find_one({"engagement_id": engagement_id}, {"_id": 0})
    if not eng:
        raise HTTPException(status_code=404, detail="Engagement not found")

    checklist = eng.get("checklist", [])
    if req.group_index < len(checklist) and req.item_index < len(checklist[req.group_index]["items"]):
        checklist[req.group_index]["items"][req.item_index]["done"] = req.done

    # Auto-update phase based on checklist progress
    phase = checklist[0]["group"] if checklist else ""
    for g in checklist:
        if any(not item["done"] for item in g["items"]):
            phase = g["group"]
            break
        phase = g["group"]

    total = sum(len(g["items"]) for g in checklist)
    done_count = sum(1 for g in checklist for item in g["items"] if item["done"])
    status = "Completed" if done_count == total and total > 0 else "Active"

    await db.service_engagements.update_one(
        {"engagement_id": engagement_id},
        {"$set": {"checklist": checklist, "phase": phase, "status": status}}
    )
    return {"phase": phase, "status": status, "progress": round(done_count / total * 100) if total > 0 else 0}

@api_router.patch("/service/engagements/{engagement_id}")
async def update_service_engagement(engagement_id: str, status: Optional[str] = None, phase: Optional[str] = None, assigned_to: Optional[str] = None, assigned_to_name: Optional[str] = None, authorization: str = Header(None), session_token: str = Cookie(None)):
    user = await get_current_user(authorization, session_token)
    update_data = {}
    if status:
        update_data["status"] = status
    if phase:
        update_data["phase"] = phase
    if assigned_to:
        update_data["assigned_to"] = assigned_to
    if assigned_to_name:
        update_data["assigned_to_name"] = assigned_to_name
    if not update_data:
        raise HTTPException(status_code=400, detail="Nothing to update")
    await db.service_engagements.update_one({"engagement_id": engagement_id}, {"$set": update_data})
    return {"message": "Engagement updated"}

# ============= BYPASS APPROVAL (Quick Approve) =============

@api_router.post("/service/engagements/{engagement_id}/approve")
async def bypass_approve_engagement(engagement_id: str, authorization: str = Header(None), session_token: str = Cookie(None)):
    """Partner can directly approve all remaining checklist items and mark engagement as Completed."""
    user = await get_current_user(authorization, session_token)
    if user.get("role") != "partner":
        raise HTTPException(status_code=403, detail="Only partners can bypass-approve")

    eng = await db.service_engagements.find_one({"engagement_id": engagement_id}, {"_id": 0})
    if not eng:
        raise HTTPException(status_code=404, detail="Engagement not found")

    # Mark all checklist items as done
    checklist = eng.get("checklist", [])
    for group in checklist:
        for item in group.get("items", []):
            item["done"] = True

    # Set last group as current phase
    last_phase = checklist[-1]["group"] if checklist else eng.get("phase", "")

    await db.service_engagements.update_one(
        {"engagement_id": engagement_id},
        {"$set": {
            "checklist": checklist,
            "status": "Completed",
            "phase": last_phase,
            "approved_by": user["user_id"],
            "approved_by_name": user.get("name"),
            "approved_at": datetime.now(timezone.utc).isoformat(),
        }}
    )
    await log_activity("Bypass approval", f"Partner {user.get('name')} approved {eng.get('service_type')} for {eng.get('client_name')}", user["user_id"], eng.get("client_id"), eng.get("client_name"))
    return {"message": f"Engagement approved by {user.get('name')}", "status": "Completed", "phase": last_phase}

# ============= SERVICE-BASED REMINDER CONFIG =============

@api_router.get("/settings/reminder-config")
async def get_reminder_config(authorization: str = Header(None), session_token: str = Cookie(None)):
    user = await require_partner(authorization, session_token)
    doc = await db.settings.find_one({"type": "reminder_config"}, {"_id": 0})
    if not doc:
        return {"type": "reminder_config", "aml_designated_staff": "", "audit_client_mapping": {}}
    return doc

class ReminderConfigRequest(BaseModel):
    aml_designated_staff: Optional[str] = None
    audit_client_mapping: Optional[Dict[str, str]] = None

@api_router.patch("/settings/reminder-config")
async def update_reminder_config(req: ReminderConfigRequest, authorization: str = Header(None), session_token: str = Cookie(None)):
    user = await require_partner(authorization, session_token)
    update = {"type": "reminder_config", "updated_by": user["user_id"], "updated_at": datetime.now(timezone.utc).isoformat()}
    if req.aml_designated_staff is not None:
        update["aml_designated_staff"] = req.aml_designated_staff
    if req.audit_client_mapping is not None:
        update["audit_client_mapping"] = req.audit_client_mapping
    await db.settings.update_one({"type": "reminder_config"}, {"$set": update}, upsert=True)
    return {"message": "Reminder config updated"}

# ============= SETTINGS ROUTES (Partner Only) =============

SECTIONS = ["overview", "audit", "vat", "corporate", "advisory", "aml"]

async def require_partner(authorization, session_token):
    user = await get_current_user(authorization, session_token)
    if user.get("role") != "partner":
        raise HTTPException(status_code=403, detail="Partners only")
    return user

# --- RBAC ---
@api_router.get("/settings/rbac")
async def get_rbac(authorization: str = Header(None), session_token: str = Cookie(None)):
    user = await require_partner(authorization, session_token)
    doc = await db.settings.find_one({"type": "rbac"}, {"_id": 0})
    if not doc:
        default = {"staff": {s: True for s in SECTIONS}, "partner": {s: True for s in SECTIONS}}
        return {"type": "rbac", "config": default}
    return doc

class RBACUpdateRequest(BaseModel):
    config: Dict[str, Dict[str, bool]]

@api_router.patch("/settings/rbac")
async def update_rbac(req: RBACUpdateRequest, authorization: str = Header(None), session_token: str = Cookie(None)):
    user = await require_partner(authorization, session_token)
    await db.settings.update_one({"type": "rbac"}, {"$set": {"type": "rbac", "config": req.config, "updated_by": user["user_id"], "updated_at": datetime.now(timezone.utc).isoformat()}}, upsert=True)
    return {"message": "RBAC updated"}

# Public RBAC endpoint (any authenticated user can fetch to filter their sidebar)
@api_router.get("/settings/rbac-public")
async def get_rbac_public(authorization: str = Header(None), session_token: str = Cookie(None)):
    user = await get_current_user(authorization, session_token)
    doc = await db.settings.find_one({"type": "rbac"}, {"_id": 0})
    if not doc:
        default = {"staff": {s: True for s in SECTIONS}, "partner": {s: True for s in SECTIONS}}
        return {"type": "rbac", "config": default}
    return doc

# --- Firm ---
@api_router.get("/settings/firm")
async def get_firm_settings(authorization: str = Header(None), session_token: str = Cookie(None)):
    user = await require_partner(authorization, session_token)
    doc = await db.settings.find_one({"type": "firm"}, {"_id": 0})
    if not doc:
        return {"type": "firm", "firm_name": "Nair & Nelliyatt Chartered Accountants"}
    return doc

class FirmUpdateRequest(BaseModel):
    firm_name: str

@api_router.patch("/settings/firm")
async def update_firm_settings(req: FirmUpdateRequest, authorization: str = Header(None), session_token: str = Cookie(None)):
    user = await require_partner(authorization, session_token)
    await db.settings.update_one({"type": "firm"}, {"$set": {"type": "firm", "firm_name": req.firm_name, "updated_by": user["user_id"], "updated_at": datetime.now(timezone.utc).isoformat()}}, upsert=True)
    return {"message": "Firm name updated"}

# --- User Management ---
@api_router.get("/settings/users")
async def get_all_users(authorization: str = Header(None), session_token: str = Cookie(None)):
    user = await require_partner(authorization, session_token)
    users = await db.users.find({}, {"_id": 0, "password_hash": 0, "password": 0}).to_list(100)
    return users

class UpdateUserRequest(BaseModel):
    role: Optional[str] = None
    title: Optional[str] = None
    email: Optional[str] = None
    new_password: Optional[str] = None
    date_of_joining: Optional[str] = None
    client_id: Optional[str] = None

@api_router.patch("/settings/users/{user_id}")
async def update_user(user_id: str, req: UpdateUserRequest, authorization: str = Header(None), session_token: str = Cookie(None)):
    user = await require_partner(authorization, session_token)
    target = await db.users.find_one({"user_id": user_id})
    if not target:
        raise HTTPException(status_code=404, detail="User not found")
    update_data = {}
    if req.role and req.role in ("staff", "partner", "client"):
        update_data["role"] = req.role
        # Clear client_id when role changes away from client
        if req.role != "client":
            update_data["client_id"] = None
    if req.client_id is not None and req.role == "client":
        update_data["client_id"] = req.client_id
    if req.title is not None:
        update_data["title"] = req.title
    if req.email is not None and req.email != target.get("email"):
        existing = await db.users.find_one({"email": req.email, "user_id": {"$ne": user_id}})
        if existing:
            raise HTTPException(status_code=400, detail="Email already in use by another user")
        update_data["email"] = req.email
    if req.date_of_joining is not None:
        update_data["date_of_joining"] = req.date_of_joining
    if req.new_password and len(req.new_password) >= 6:
        update_data["password"] = bcrypt.hashpw(req.new_password.encode(), bcrypt.gensalt()).decode()
    if not update_data:
        raise HTTPException(status_code=400, detail="Nothing to update")
    await db.users.update_one({"user_id": user_id}, {"$set": update_data})
    await log_activity("User updated", f"Updated {target.get('name', user_id)}: {', '.join(update_data.keys())}", user["user_id"])
    return {"message": f"User {target.get('name')} updated"}

@api_router.patch("/settings/users/{user_id}/approve")
async def approve_user(user_id: str, authorization: str = Header(None), session_token: str = Cookie(None)):
    user = await require_partner(authorization, session_token)
    target = await db.users.find_one({"user_id": user_id})
    if not target:
        raise HTTPException(status_code=404, detail="User not found")
    await db.users.update_one({"user_id": user_id}, {"$set": {"status": "approved"}})
    await log_activity("User approved", f"Approved {target.get('name', user_id)} ({target.get('email')})", user["user_id"])
    return {"message": f"User {target.get('name')} approved"}

@api_router.patch("/settings/users/{user_id}/reject")
async def reject_user(user_id: str, authorization: str = Header(None), session_token: str = Cookie(None)):
    user = await require_partner(authorization, session_token)
    target = await db.users.find_one({"user_id": user_id})
    if not target:
        raise HTTPException(status_code=404, detail="User not found")
    await db.users.delete_one({"user_id": user_id})
    await db.user_sessions.delete_many({"user_id": user_id})
    await log_activity("User rejected", f"Rejected {target.get('name', user_id)} ({target.get('email')})", user["user_id"])
    return {"message": f"User {target.get('name')} rejected and removed"}

# --- Storage ---
@api_router.get("/settings/storage")
async def get_storage_config(authorization: str = Header(None), session_token: str = Cookie(None)):
    user = await require_partner(authorization, session_token)
    doc = await db.settings.find_one({"type": "storage"}, {"_id": 0})
    if not doc:
        return {"type": "storage", "config": {"provider": "default", "aws_s3": {}, "google_drive": {}, "onedrive": {}}}
    return doc

class StorageUpdateRequest(BaseModel):
    config: Dict[str, Any]

@api_router.patch("/settings/storage")
async def update_storage_config(req: StorageUpdateRequest, authorization: str = Header(None), session_token: str = Cookie(None)):
    user = await require_partner(authorization, session_token)
    await db.settings.update_one({"type": "storage"}, {"$set": {"type": "storage", "config": req.config, "updated_by": user["user_id"], "updated_at": datetime.now(timezone.utc).isoformat()}}, upsert=True)
    return {"message": "Storage config updated"}

# --- Workflows ---
@api_router.get("/settings/workflows")
async def get_workflows(authorization: str = Header(None), session_token: str = Cookie(None), service_type: Optional[str] = None):
    user = await require_partner(authorization, session_token)
    query = {"is_deleted": {"$ne": True}}
    if service_type:
        query["service_type"] = service_type
    wfs = await db.workflows.find(query, {"_id": 0}).sort("created_at", 1).to_list(200)
    return wfs

class WorkflowStep(BaseModel):
    name: str
    description: Optional[str] = ""
    order: int

class CreateWorkflowRequest(BaseModel):
    name: str
    service_type: str
    steps: List[WorkflowStep]
    is_preset: bool = False

@api_router.post("/settings/workflows")
async def create_workflow(req: CreateWorkflowRequest, authorization: str = Header(None), session_token: str = Cookie(None)):
    user = await require_partner(authorization, session_token)
    wf_id = f"wf_{uuid.uuid4().hex[:12]}"
    wf_doc = {
        "workflow_id": wf_id,
        "name": req.name,
        "service_type": req.service_type,
        "steps": [s.model_dump() for s in req.steps],
        "is_preset": req.is_preset,
        "created_by": user["user_id"],
        "created_by_name": user.get("name"),
        "created_at": datetime.now(timezone.utc).isoformat(),
    }
    await db.workflows.insert_one(wf_doc)
    wf_doc.pop("_id", None)
    return wf_doc

class UpdateWorkflowRequest(BaseModel):
    name: Optional[str] = None
    steps: Optional[List[WorkflowStep]] = None

@api_router.patch("/settings/workflows/{workflow_id}")
async def update_workflow(workflow_id: str, req: UpdateWorkflowRequest, authorization: str = Header(None), session_token: str = Cookie(None)):
    user = await require_partner(authorization, session_token)
    update = {}
    if req.name:
        update["name"] = req.name
    if req.steps is not None:
        update["steps"] = [s.model_dump() for s in req.steps]
    if not update:
        raise HTTPException(status_code=400, detail="Nothing to update")
    update["updated_at"] = datetime.now(timezone.utc).isoformat()
    result = await db.workflows.update_one({"workflow_id": workflow_id}, {"$set": update})
    if result.matched_count == 0:
        raise HTTPException(status_code=404, detail="Workflow not found")
    return {"message": "Workflow updated"}

@api_router.delete("/settings/workflows/{workflow_id}")
async def delete_workflow(workflow_id: str, authorization: str = Header(None), session_token: str = Cookie(None)):
    user = await require_partner(authorization, session_token)
    result = await db.workflows.update_one({"workflow_id": workflow_id}, {"$set": {"is_deleted": True}})
    if result.matched_count == 0:
        raise HTTPException(status_code=404, detail="Workflow not found")
    return {"message": "Workflow deleted"}

# ============= NOTIFICATIONS =============

@api_router.get("/notifications")
async def get_notifications(authorization: str = Header(None), session_token: str = Cookie(None)):
    user = await get_current_user(authorization, session_token)
    today = datetime.now(timezone.utc).date()
    notifications = []

    # Check overdue & upcoming tasks assigned to user (or all if partner)
    task_query = {"status": {"$ne": "Completed"}}
    if user.get("role") != "partner":
        task_query["assigned_to"] = user.get("email")

    tasks = await db.tasks.find(task_query, {"_id": 0}).to_list(100)
    for task in tasks:
        due = task.get("due_date")
        if not due:
            continue
        try:
            due_date = datetime.fromisoformat(due).date() if 'T' in due else datetime.strptime(due, "%Y-%m-%d").date()
        except (ValueError, TypeError):
            continue

        days_until = (due_date - today).days
        if days_until < 0:
            notifications.append({
                "notification_id": f"notif_task_{task.get('task_id', '')}",
                "title": f"Overdue: {task.get('title', 'Task')}",
                "message": f"Was due {abs(days_until)} day{'s' if abs(days_until) != 1 else ''} ago — {task.get('client_name', '')}",
                "severity": "urgent",
                "due_date": due,
                "type": "overdue_task",
                "ref_id": task.get("task_id"),
            })
        elif days_until <= 3:
            notifications.append({
                "notification_id": f"notif_task_{task.get('task_id', '')}",
                "title": f"Due Soon: {task.get('title', 'Task')}",
                "message": f"Due in {days_until} day{'s' if days_until != 1 else ''} — {task.get('client_name', '')}",
                "severity": "warning" if days_until <= 1 else "info",
                "due_date": due,
                "type": "upcoming_task",
                "ref_id": task.get("task_id"),
            })

    # Check upcoming events in next 3 days
    events = await db.events.find({"date": {"$exists": True}}, {"_id": 0}).to_list(100)
    for ev in events:
        ev_date_str = ev.get("date")
        if not ev_date_str:
            continue
        try:
            ev_date = datetime.fromisoformat(ev_date_str).date() if 'T' in ev_date_str else datetime.strptime(ev_date_str, "%Y-%m-%d").date()
        except (ValueError, TypeError):
            continue

        days_until = (ev_date - today).days
        if 0 <= days_until <= 3:
            day_label = "Today" if days_until == 0 else f"In {days_until} day{'s' if days_until != 1 else ''}"
            notifications.append({
                "notification_id": f"notif_ev_{ev.get('event_id', '')}",
                "title": f"Upcoming: {ev.get('title', 'Event')}",
                "message": f"{day_label} — {ev.get('event_type', 'event').capitalize()}",
                "severity": "warning" if days_until == 0 else "info",
                "due_date": ev_date_str,
                "type": "upcoming_event",
                "ref_id": ev.get("event_id"),
            })

    # Check dismissed notifications
    dismissed = await db.dismissed_notifications.find({"user_id": user["user_id"]}, {"_id": 0, "notification_id": 1}).to_list(500)
    dismissed_ids = {d["notification_id"] for d in dismissed}
    notifications = [n for n in notifications if n["notification_id"] not in dismissed_ids]

    # Sort: urgent first, then by due_date
    severity_order = {"urgent": 0, "warning": 1, "info": 2}
    notifications.sort(key=lambda x: (severity_order.get(x["severity"], 2), x.get("due_date", "")))
    return notifications[:20]

@api_router.patch("/notifications/{notification_id}/dismiss")
async def dismiss_notification(notification_id: str, authorization: str = Header(None), session_token: str = Cookie(None)):
    user = await get_current_user(authorization, session_token)
    await db.dismissed_notifications.update_one(
        {"user_id": user["user_id"], "notification_id": notification_id},
        {"$set": {"user_id": user["user_id"], "notification_id": notification_id, "dismissed_at": datetime.now(timezone.utc).isoformat()}},
        upsert=True
    )
    return {"message": "Notification dismissed"}

# ============= WORKFLOW-LINKED ENGAGEMENTS =============

@api_router.get("/service/workflows-for-type")
async def get_workflows_for_service(service_type: str, authorization: str = Header(None), session_token: str = Cookie(None)):
    user = await get_current_user(authorization, session_token)
    wfs = await db.workflows.find({"service_type": service_type, "is_deleted": {"$ne": True}}, {"_id": 0}).to_list(20)
    return wfs

# ============= CLIENT ACTIVITY TIMELINE =============

@api_router.get("/clients/{client_id}/timeline")
async def get_client_timeline(client_id: str, authorization: str = Header(None), session_token: str = Cookie(None)):
    user = await get_current_user(authorization, session_token)
    client = await db.clients.find_one({"client_id": client_id}, {"_id": 0})
    if not client:
        raise HTTPException(status_code=404, detail="Client not found")

    timeline = []

    # Tasks
    tasks = await db.tasks.find({"client_id": client_id}, {"_id": 0}).to_list(100)
    for t in tasks:
        timeline.append({
            "type": "task", "title": t.get("title", "Task"),
            "status": t.get("status"), "date": t.get("due_date") or t.get("created_at", ""),
            "detail": f"Assigned to {t.get('assigned_to_name', t.get('assigned_to', '—'))} | Priority: {t.get('priority', '—')}",
            "ref_id": t.get("task_id"),
        })

    # Events
    events = await db.events.find({"client_id": client_id}, {"_id": 0}).to_list(100)
    for e in events:
        timeline.append({
            "type": e.get("event_type", "event"), "title": e.get("title", "Event"),
            "status": "Scheduled", "date": e.get("date", ""),
            "detail": e.get("description", ""),
            "ref_id": e.get("event_id"),
        })

    # Documents
    docs = await db.documents.find({"client_id": client_id, "is_deleted": False}, {"_id": 0}).to_list(100)
    for d in docs:
        timeline.append({
            "type": "document", "title": d.get("original_filename", "Document"),
            "status": d.get("document_type", ""), "date": d.get("created_at", ""),
            "detail": f"Uploaded by {d.get('uploaded_by_name', '—')}",
            "ref_id": d.get("file_id"),
        })

    # Engagements
    engs = await db.service_engagements.find({"client_id": client_id, "status": {"$ne": "Deleted"}}, {"_id": 0}).to_list(50)
    for eng in engs:
        total = sum(len(g.get("items", [])) for g in eng.get("checklist", []))
        done = sum(1 for g in eng.get("checklist", []) for item in g.get("items", []) if item.get("done"))
        progress = round(done / total * 100) if total > 0 else 0
        timeline.append({
            "type": "engagement", "title": eng.get("service_type", "").replace("_", " ").title(),
            "status": eng.get("status"), "date": eng.get("created_at", ""),
            "detail": f"Phase: {eng.get('phase', '—')} | Progress: {progress}% | {eng.get('assigned_to_name', '—')}",
            "ref_id": eng.get("engagement_id"),
        })

    # Activities
    activities = await db.activities.find({"client_id": client_id}, {"_id": 0}).sort("timestamp", -1).to_list(50)
    for a in activities:
        timeline.append({
            "type": "activity", "title": a.get("action", "Activity"),
            "status": "", "date": a.get("timestamp", ""),
            "detail": a.get("description", ""),
            "ref_id": None,
        })

    # Sort by date descending
    def parse_date(item):
        d = item.get("date", "")
        if not d:
            return ""
        return d

    timeline.sort(key=parse_date, reverse=True)
    return {"client": client, "timeline": timeline}

# ============= CLIENT ONBOARDING =============

ONBOARDING_DOCUMENT_CHECKLIST = {
    "general": ["Trade Licence Copy", "Memorandum of Association", "Emirates ID (Partners/Directors)", "Passport Copies (Partners/Directors)", "Power of Attorney (if applicable)"],
    "vat": ["Bank Letter / IBAN Certificate", "Turnover Evidence (12 months)", "Previous VAT Returns (if any)"],
    "corporate_tax": ["Financial Statements (Latest)", "Trial Balance", "Tax Registration Certificate (if any)"],
    "aml": ["Source of Funds Declaration", "Beneficial Ownership Structure", "Sanctions Self-Declaration"],
}

SERVICE_TO_ENGAGEMENT_TYPE = {
    "Statutory Audit": "statutory_audit",
    "Internal Audit": "internal_audit",
    "Stock Audit": "stock_audit",
    "Fraud Audit": "fraud_audit",
    "VAT Registration": "vat_registration",
    "VAT Filing": "vat_filing",
    "VAT Amendments": "vat_filing",
    "Corporate Registration": "corporate_tax",
    "Corporate Tax": "corporate_tax",
    "Company Formation": "corporate_tax",
    "Liquidation": "corporate_tax",
    "Valuation": "corporate_tax",
    "Due Diligence": "corporate_tax",
    "AML Review": "aml_review",
    "AML Filing": "aml_review",
    "AML Report": "aml_review",
}

class OnboardingRequest(BaseModel):
    name: str
    entity_type: str
    jurisdiction: Optional[str] = None
    trade_licence_no: Optional[str] = None
    trn: Optional[str] = None
    ct_registration_no: Optional[str] = None
    aml_risk_rating: str = "Low"
    pep_flag: bool = False
    active_services: List[str] = []
    relationship_manager: Optional[str] = None
    relationship_manager_name: Optional[str] = None
    contact_person: Optional[str] = None
    contact_email: Optional[str] = None
    contact_phone: Optional[str] = None
    notes: Optional[str] = None
    auto_create_engagements: bool = True

@api_router.post("/onboarding")
async def onboard_client(req: OnboardingRequest, authorization: str = Header(None), session_token: str = Cookie(None)):
    user = await get_current_user(authorization, session_token)
    if user.get("role") != "partner":
        raise HTTPException(status_code=403, detail="Only partners can onboard clients")

    # Check if client already exists
    existing = await db.clients.find_one({"name": req.name})
    if existing:
        raise HTTPException(status_code=400, detail=f"Client '{req.name}' already exists")

    # 1. Create Client
    client_id = f"client_{uuid.uuid4().hex[:12]}"
    client_doc = {
        "client_id": client_id,
        "name": req.name,
        "entity_type": req.entity_type,
        "jurisdiction": req.jurisdiction,
        "trade_licence_no": req.trade_licence_no,
        "trn": req.trn,
        "ct_registration_no": req.ct_registration_no,
        "aml_risk_rating": req.aml_risk_rating,
        "pep_flag": req.pep_flag,
        "status": "Onboarding",
        "active_services": req.active_services,
        "relationship_manager_id": req.relationship_manager,
        "contact_person": req.contact_person,
        "contact_email": req.contact_email,
        "contact_phone": req.contact_phone,
        "notes": req.notes,
        "created_at": datetime.now(timezone.utc).isoformat(),
        "onboarded_by": user["user_id"],
        "onboarded_by_name": user.get("name"),
    }
    await db.clients.insert_one(client_doc)

    # 2. Generate onboarding tasks
    tasks_created = []
    assignee = req.relationship_manager or user.get("email")
    assignee_name = req.relationship_manager_name or user.get("name")
    due_base = datetime.now(timezone.utc) + timedelta(days=3)

    # General KYC tasks
    onboarding_tasks = [
        {"title": f"Collect KYC Documents - {req.name}", "service_module": "Onboarding", "priority": "High", "days_offset": 3},
        {"title": f"Verify Entity Details - {req.name}", "service_module": "Onboarding", "priority": "Medium", "days_offset": 5},
        {"title": f"Complete AML Risk Assessment - {req.name}", "service_module": "AML", "priority": "High", "days_offset": 5},
    ]

    # Service-specific tasks
    has_vat = any(s for s in req.active_services if "VAT" in s)
    has_audit = any(s for s in req.active_services if "Audit" in s)
    has_ct = any(s for s in req.active_services if "Corporate" in s or "Tax" in s)

    if has_vat:
        onboarding_tasks.append({"title": f"VAT Registration Check - {req.name}", "service_module": "VAT", "priority": "Medium", "days_offset": 7})
    if has_audit:
        onboarding_tasks.append({"title": f"Schedule Audit Kickoff Meeting - {req.name}", "service_module": "Audit", "priority": "Medium", "days_offset": 10})
    if has_ct:
        onboarding_tasks.append({"title": f"Corporate Tax Registration Review - {req.name}", "service_module": "Corporate Tax", "priority": "Medium", "days_offset": 7})

    for t in onboarding_tasks:
        task_id = f"task_{uuid.uuid4().hex[:12]}"
        task_doc = {
            "task_id": task_id,
            "title": t["title"],
            "description": f"Auto-generated onboarding task for {req.name}",
            "service_module": t["service_module"],
            "client_id": client_id,
            "client_name": req.name,
            "due_date": (datetime.now(timezone.utc) + timedelta(days=t["days_offset"])).strftime("%Y-%m-%d"),
            "priority": t["priority"],
            "assigned_to": assignee,
            "assigned_to_name": assignee_name,
            "status": "Pending",
            "created_by": user["user_id"],
            "created_at": datetime.now(timezone.utc).isoformat(),
        }
        await db.tasks.insert_one(task_doc)
        tasks_created.append(task_id)

    # 3. Auto-create engagements for selected services (if enabled)
    engagements_created = []
    if req.auto_create_engagements and req.active_services:
        # Deduplicate engagement types
        engagement_types_done = set()
        for svc in req.active_services:
            eng_type = SERVICE_TO_ENGAGEMENT_TYPE.get(svc)
            if eng_type and eng_type not in engagement_types_done:
                engagement_types_done.add(eng_type)

                # Check for a workflow template for this service type
                wf = await db.workflows.find_one({"service_type": eng_type, "is_deleted": {"$ne": True}}, {"_id": 0})
                checklist = []
                workflow_name = None
                if wf:
                    workflow_name = wf.get("name")
                    checklist = [{"group": step["name"], "items": [{"label": step.get("description") or step["name"], "done": False}]} for step in wf.get("steps", [])]
                if not checklist:
                    template = CHECKLIST_TEMPLATES.get(eng_type, [])
                    checklist = [{"group": g["group"], "items": [{"label": item, "done": False} for item in g["items"]]} for g in template]

                eng_id = f"eng_{uuid.uuid4().hex[:12]}"
                eng_doc = {
                    "engagement_id": eng_id,
                    "service_type": eng_type,
                    "client_id": client_id,
                    "client_name": req.name,
                    "assigned_to": assignee,
                    "assigned_to_name": assignee_name,
                    "status": "Active",
                    "phase": checklist[0]["group"] if checklist else "",
                    "checklist": checklist,
                    "workflow_id": wf.get("workflow_id") if wf else None,
                    "workflow_name": workflow_name,
                    "created_by": user["user_id"],
                    "created_by_name": user.get("name"),
                    "created_at": datetime.now(timezone.utc).isoformat(),
                }
                await db.service_engagements.insert_one(eng_doc)
                engagements_created.append(eng_id)

    # 4. Generate document checklist items
    doc_checklist = list(ONBOARDING_DOCUMENT_CHECKLIST["general"])
    if has_vat:
        doc_checklist.extend(ONBOARDING_DOCUMENT_CHECKLIST["vat"])
    if has_ct:
        doc_checklist.extend(ONBOARDING_DOCUMENT_CHECKLIST["corporate_tax"])
    if req.aml_risk_rating in ("Medium", "High"):
        doc_checklist.extend(ONBOARDING_DOCUMENT_CHECKLIST["aml"])

    # 5. Log activity
    await log_activity("Client onboarded", f"Onboarded new client: {req.name} ({req.entity_type})", user["user_id"], client_id, req.name)

    return {
        "client_id": client_id,
        "client_name": req.name,
        "status": "Onboarding",
        "tasks_created": len(tasks_created),
        "engagements_created": len(engagements_created),
        "document_checklist": doc_checklist,
        "message": f"Client '{req.name}' onboarded successfully",
    }

@api_router.get("/onboarding/document-checklist")
async def get_onboarding_doc_checklist(services: str = "", risk: str = "Low", authorization: str = Header(None), session_token: str = Cookie(None)):
    """Preview document checklist based on selected services and risk."""
    user = await get_current_user(authorization, session_token)
    svc_list = [s.strip() for s in services.split(",") if s.strip()] if services else []

    checklist = list(ONBOARDING_DOCUMENT_CHECKLIST["general"])
    has_vat = any(s for s in svc_list if "VAT" in s)
    has_ct = any(s for s in svc_list if "Corporate" in s or "Tax" in s)
    if has_vat:
        checklist.extend(ONBOARDING_DOCUMENT_CHECKLIST["vat"])
    if has_ct:
        checklist.extend(ONBOARDING_DOCUMENT_CHECKLIST["corporate_tax"])
    if risk in ("Medium", "High"):
        checklist.extend(ONBOARDING_DOCUMENT_CHECKLIST["aml"])

    return {"checklist": checklist}

# ============= EXPORT / REPORTING =============

@api_router.get("/export/audit-report/{engagement_id}")
async def export_audit_report(engagement_id: str, authorization: str = Header(None), session_token: str = Cookie(None)):
    from fpdf import FPDF
    import io

    user = await get_current_user(authorization, session_token)
    eng = await db.service_engagements.find_one({"engagement_id": engagement_id}, {"_id": 0})
    if not eng:
        raise HTTPException(status_code=404, detail="Engagement not found")

    # Get firm name
    firm_doc = await db.settings.find_one({"type": "firm"}, {"_id": 0})
    firm_name = firm_doc.get("firm_name", "Nair & Nelliyatt Chartered Accountants") if firm_doc else "Nair & Nelliyatt Chartered Accountants"

    pdf = FPDF()
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=20)

    # Header
    pdf.set_font("Helvetica", "B", 18)
    pdf.set_text_color(10, 17, 40)
    pdf.cell(0, 12, firm_name, new_x="LMARGIN", new_y="NEXT", align="C")
    pdf.set_font("Helvetica", "", 10)
    pdf.set_text_color(120, 120, 120)
    pdf.cell(0, 6, "Engagement Report", new_x="LMARGIN", new_y="NEXT", align="C")
    pdf.ln(8)

    # Engagement Info
    pdf.set_draw_color(212, 175, 55)
    pdf.set_line_width(0.5)
    pdf.line(10, pdf.get_y(), 200, pdf.get_y())
    pdf.ln(6)

    pdf.set_font("Helvetica", "B", 12)
    pdf.set_text_color(10, 17, 40)
    pdf.cell(0, 8, eng.get("service_type", "").replace("_", " ").title(), new_x="LMARGIN", new_y="NEXT")

    info_lines = [
        f"Client: {eng.get('client_name', '—')}",
        f"Assigned To: {eng.get('assigned_to_name', '—')}",
        f"Status: {eng.get('status', '—')}",
        f"Current Phase: {eng.get('phase', '—')}",
        f"Created: {eng.get('created_at', '—')[:10]}",
    ]
    if eng.get("workflow_name"):
        info_lines.append(f"Workflow: {eng.get('workflow_name')}")

    pdf.set_font("Helvetica", "", 10)
    pdf.set_text_color(60, 60, 60)
    for line in info_lines:
        pdf.cell(0, 6, line, new_x="LMARGIN", new_y="NEXT")
    pdf.ln(6)

    # Checklist Progress
    total = sum(len(g.get("items", [])) for g in eng.get("checklist", []))
    done = sum(1 for g in eng.get("checklist", []) for item in g.get("items", []) if item.get("done"))
    progress = round(done / total * 100) if total > 0 else 0

    pdf.set_font("Helvetica", "B", 11)
    pdf.set_text_color(10, 17, 40)
    pdf.cell(0, 8, f"Checklist Progress: {done}/{total} ({progress}%)", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)

    for group in eng.get("checklist", []):
        pdf.set_font("Helvetica", "B", 10)
        pdf.set_text_color(212, 175, 55)
        pdf.cell(0, 7, group.get("group", ""), new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "", 9)
        pdf.set_text_color(60, 60, 60)
        for item in group.get("items", []):
            mark = "[x]" if item.get("done") else "[ ]"
            pdf.cell(0, 5, f"  {mark} {item.get('label', '')}", new_x="LMARGIN", new_y="NEXT")
        pdf.ln(2)

    # Footer
    pdf.ln(10)
    pdf.set_font("Helvetica", "I", 8)
    pdf.set_text_color(150, 150, 150)
    pdf.cell(0, 5, f"Generated on {datetime.now(timezone.utc).strftime('%d %b %Y at %H:%M UTC')} by {user.get('name', '—')}", new_x="LMARGIN", new_y="NEXT", align="C")

    buffer = io.BytesIO()
    pdf.output(buffer)
    buffer.seek(0)

    filename = f"Audit_Report_{eng.get('client_name', 'client').replace(' ', '_')}_{engagement_id[-6:]}.pdf"
    return StreamingResponse(buffer, media_type="application/pdf", headers={"Content-Disposition": f"attachment; filename={filename}"})

@api_router.get("/export/vat-return/{engagement_id}")
async def export_vat_return(engagement_id: str, authorization: str = Header(None), session_token: str = Cookie(None)):
    from fpdf import FPDF
    import io

    user = await get_current_user(authorization, session_token)
    eng = await db.service_engagements.find_one({"engagement_id": engagement_id}, {"_id": 0})
    if not eng:
        raise HTTPException(status_code=404, detail="Engagement not found")

    firm_doc = await db.settings.find_one({"type": "firm"}, {"_id": 0})
    firm_name = firm_doc.get("firm_name", "Nair & Nelliyatt Chartered Accountants") if firm_doc else "Nair & Nelliyatt Chartered Accountants"

    pdf = FPDF()
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=20)

    # Header
    pdf.set_font("Helvetica", "B", 18)
    pdf.set_text_color(10, 17, 40)
    pdf.cell(0, 12, firm_name, new_x="LMARGIN", new_y="NEXT", align="C")
    pdf.set_font("Helvetica", "", 10)
    pdf.set_text_color(120, 120, 120)
    pdf.cell(0, 6, "VAT Return Summary", new_x="LMARGIN", new_y="NEXT", align="C")
    pdf.ln(8)

    pdf.set_draw_color(212, 175, 55)
    pdf.set_line_width(0.5)
    pdf.line(10, pdf.get_y(), 200, pdf.get_y())
    pdf.ln(6)

    # Client info
    pdf.set_font("Helvetica", "B", 12)
    pdf.set_text_color(10, 17, 40)
    pdf.cell(0, 8, f"Client: {eng.get('client_name', '—')}", new_x="LMARGIN", new_y="NEXT")

    pdf.set_font("Helvetica", "", 10)
    pdf.set_text_color(60, 60, 60)
    pdf.cell(0, 6, f"Filing Status: {eng.get('status', '—')}", new_x="LMARGIN", new_y="NEXT")
    pdf.cell(0, 6, f"Current Phase: {eng.get('phase', '—')}", new_x="LMARGIN", new_y="NEXT")
    pdf.cell(0, 6, f"Prepared By: {eng.get('assigned_to_name', '—')}", new_x="LMARGIN", new_y="NEXT")
    pdf.cell(0, 6, f"Period: {eng.get('created_at', '—')[:10]}", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(6)

    # Checklist as filing steps
    pdf.set_font("Helvetica", "B", 11)
    pdf.set_text_color(10, 17, 40)
    total = sum(len(g.get("items", [])) for g in eng.get("checklist", []))
    done = sum(1 for g in eng.get("checklist", []) for item in g.get("items", []) if item.get("done"))
    pdf.cell(0, 8, f"Filing Checklist: {done}/{total} steps completed", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)

    for group in eng.get("checklist", []):
        pdf.set_font("Helvetica", "B", 10)
        pdf.set_text_color(212, 175, 55)
        pdf.cell(0, 7, group.get("group", ""), new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "", 9)
        pdf.set_text_color(60, 60, 60)
        for item in group.get("items", []):
            status = "Completed" if item.get("done") else "Pending"
            pdf.cell(0, 5, f"  [{status}] {item.get('label', '')}", new_x="LMARGIN", new_y="NEXT")
        pdf.ln(2)

    # Footer
    pdf.ln(10)
    pdf.set_font("Helvetica", "I", 8)
    pdf.set_text_color(150, 150, 150)
    pdf.cell(0, 5, f"Generated on {datetime.now(timezone.utc).strftime('%d %b %Y at %H:%M UTC')} by {user.get('name', '—')}", new_x="LMARGIN", new_y="NEXT", align="C")

    buffer = io.BytesIO()
    pdf.output(buffer)
    buffer.seek(0)

    filename = f"VAT_Return_{eng.get('client_name', 'client').replace(' ', '_')}_{engagement_id[-6:]}.pdf"
    return StreamingResponse(buffer, media_type="application/pdf", headers={"Content-Disposition": f"attachment; filename={filename}"})

# ============= BILLABLE HOURS =============

class LogHoursRequest(BaseModel):
    client_id: str
    client_name: Optional[str] = None
    task_id: Optional[str] = None
    task_title: Optional[str] = None
    hours: float
    date: str
    description: Optional[str] = None

@api_router.post("/billable-hours")
async def log_billable_hours(req: LogHoursRequest, authorization: str = Header(None), session_token: str = Cookie(None)):
    user = await get_current_user(authorization, session_token)
    entry_id = f"bh_{uuid.uuid4().hex[:12]}"

    # Auto-resolve client name if not provided
    client_name = req.client_name
    if not client_name:
        client_doc = await db.clients.find_one({"client_id": req.client_id}, {"_id": 0, "name": 1})
        client_name = client_doc["name"] if client_doc else "Unknown"

    entry = {
        "entry_id": entry_id,
        "staff_id": user["user_id"],
        "staff_name": user.get("name"),
        "staff_email": user.get("email"),
        "client_id": req.client_id,
        "client_name": client_name,
        "task_id": req.task_id,
        "task_title": req.task_title,
        "hours": req.hours,
        "date": req.date,
        "description": req.description,
        "created_at": datetime.now(timezone.utc).isoformat(),
    }
    await db.billable_hours.insert_one(entry)
    entry.pop("_id", None)
    return entry

@api_router.get("/billable-hours")
async def get_billable_hours(authorization: str = Header(None), session_token: str = Cookie(None), staff_email: Optional[str] = None, client_id: Optional[str] = None, date_from: Optional[str] = None, date_to: Optional[str] = None):
    user = await get_current_user(authorization, session_token)
    query = {}
    # Staff can only see their own hours; partners see all
    if user.get("role") != "partner":
        query["staff_email"] = user.get("email")
    elif staff_email:
        query["staff_email"] = staff_email
    if client_id:
        query["client_id"] = client_id
    if date_from or date_to:
        date_q = {}
        if date_from:
            date_q["$gte"] = date_from
        if date_to:
            date_q["$lte"] = date_to
        if date_q:
            query["date"] = date_q

    entries = await db.billable_hours.find(query, {"_id": 0}).sort("date", -1).to_list(2000)
    return entries

@api_router.delete("/billable-hours/{entry_id}")
async def delete_billable_hours(entry_id: str, authorization: str = Header(None), session_token: str = Cookie(None)):
    user = await get_current_user(authorization, session_token)
    entry = await db.billable_hours.find_one({"entry_id": entry_id}, {"_id": 0})
    if not entry:
        raise HTTPException(status_code=404, detail="Entry not found")
    # Staff can only delete their own; partners can delete any
    if user.get("role") != "partner" and entry.get("staff_email") != user.get("email"):
        raise HTTPException(status_code=403, detail="Cannot delete another user's entry")
    await db.billable_hours.delete_one({"entry_id": entry_id})
    return {"message": "Entry deleted"}

@api_router.get("/billable-hours/summary")
async def get_billable_hours_summary(authorization: str = Header(None), session_token: str = Cookie(None), date_from: Optional[str] = None, date_to: Optional[str] = None):
    """Returns summary grouped by staff and client for reporting."""
    user = await get_current_user(authorization, session_token)
    if user.get("role") != "partner":
        raise HTTPException(status_code=403, detail="Partners only")

    query = {}
    if date_from or date_to:
        date_q = {}
        if date_from:
            date_q["$gte"] = date_from
        if date_to:
            date_q["$lte"] = date_to
        if date_q:
            query["date"] = date_q

    entries = await db.billable_hours.find(query, {"_id": 0}).to_list(5000)

    # Group by staff
    staff_summary = {}
    client_summary = {}
    for e in entries:
        staff_key = e.get("staff_email", "unknown")
        staff_name = e.get("staff_name", "Unknown")
        client_key = e.get("client_id", "unknown")
        client_name = e.get("client_name", "Unknown")
        hours = e.get("hours", 0)

        if staff_key not in staff_summary:
            staff_summary[staff_key] = {"staff_name": staff_name, "staff_email": staff_key, "total_hours": 0, "clients": {}}
        staff_summary[staff_key]["total_hours"] += hours
        if client_key not in staff_summary[staff_key]["clients"]:
            staff_summary[staff_key]["clients"][client_key] = {"client_name": client_name, "hours": 0}
        staff_summary[staff_key]["clients"][client_key]["hours"] += hours

        if client_key not in client_summary:
            client_summary[client_key] = {"client_name": client_name, "total_hours": 0, "staff": {}}
        client_summary[client_key]["total_hours"] += hours
        if staff_key not in client_summary[client_key]["staff"]:
            client_summary[client_key]["staff"][staff_key] = {"staff_name": staff_name, "hours": 0}
        client_summary[client_key]["staff"][staff_key]["hours"] += hours

    # Convert clients dict to list for JSON
    for sk in staff_summary:
        staff_summary[sk]["clients"] = list(staff_summary[sk]["clients"].values())
    for ck in client_summary:
        client_summary[ck]["staff"] = list(client_summary[ck]["staff"].values())

    return {
        "total_hours": sum(e.get("hours", 0) for e in entries),
        "total_entries": len(entries),
        "by_staff": list(staff_summary.values()),
        "by_client": list(client_summary.values()),
    }

@api_router.get("/billable-hours/export")
async def export_billable_hours(authorization: str = Header(None), session_token: str = Cookie(None), staff_email: Optional[str] = None, date_from: Optional[str] = None, date_to: Optional[str] = None):
    """Export billable hours as CSV."""
    user = await get_current_user(authorization, session_token)
    if user.get("role") != "partner":
        raise HTTPException(status_code=403, detail="Partners only")

    import io, csv

    query = {}
    if staff_email:
        query["staff_email"] = staff_email
    if date_from or date_to:
        date_q = {}
        if date_from:
            date_q["$gte"] = date_from
        if date_to:
            date_q["$lte"] = date_to
        if date_q:
            query["date"] = date_q

    entries = await db.billable_hours.find(query, {"_id": 0}).sort("date", -1).to_list(5000)

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["Date", "Staff", "Email", "Client", "Task", "Hours", "Description"])
    for e in entries:
        writer.writerow([e.get("date"), e.get("staff_name"), e.get("staff_email"), e.get("client_name"), e.get("task_title", ""), e.get("hours"), e.get("description", "")])

    content = output.getvalue()
    staff_label = staff_email.split("@")[0] if staff_email else "team"
    filename = f"Billable_Hours_{staff_label}_{datetime.now(timezone.utc).strftime('%Y%m%d')}.csv"
    return Response(content=content, media_type="text/csv", headers={"Content-Disposition": f"attachment; filename={filename}"})

# ============= DATA-DRIVEN REMINDERS =============

@api_router.get("/reminders")
async def get_reminders(authorization: str = Header(None), session_token: str = Cookie(None)):
    """Returns reminders from real tasks, engagements, and events. Applies service-based routing."""
    user = await get_current_user(authorization, session_token)
    today = datetime.now(timezone.utc).date()
    reminders = []

    # Load service-based reminder config
    reminder_cfg = await db.settings.find_one({"type": "reminder_config"}, {"_id": 0})
    aml_staff = reminder_cfg.get("aml_designated_staff", "") if reminder_cfg else ""
    audit_mapping = reminder_cfg.get("audit_client_mapping", {}) if reminder_cfg else {}

    # Overdue & upcoming tasks
    task_query = {"status": {"$ne": "Completed"}}
    if user.get("role") != "partner":
        # Staff see their own + service-routed tasks
        user_email = user.get("email")
        # If this staff is the designated AML person, also show AML tasks
        service_or = [{"assigned_to": user_email}]
        if aml_staff == user_email:
            service_or.append({"service_module": {"$regex": "AML", "$options": "i"}})
        # If this staff is mapped for any audit client, show those too
        audit_clients = [cid for cid, staff in audit_mapping.items() if staff == user_email]
        if audit_clients:
            service_or.append({"client_id": {"$in": audit_clients}, "service_module": {"$regex": "Audit", "$options": "i"}})
        task_query["$or"] = service_or

    tasks = await db.tasks.find(task_query, {"_id": 0}).to_list(200)
    for t in tasks:
        due = t.get("due_date")
        if not due:
            continue
        try:
            due_date = datetime.fromisoformat(due).date() if 'T' in due else datetime.strptime(due, "%Y-%m-%d").date()
        except (ValueError, TypeError):
            continue
        days_until = (due_date - today).days
        if days_until <= 14:
            reminders.append({
                "reminder_id": f"rem_task_{t.get('task_id', '')}",
                "type": "task",
                "title": t.get("title", "Task"),
                "client_name": t.get("client_name", ""),
                "client_id": t.get("client_id"),
                "assigned_to": t.get("assigned_to_name") or t.get("assigned_to", ""),
                "due_date": due,
                "days_until": days_until,
                "priority": t.get("priority", "Medium"),
                "status": t.get("status", "Pending"),
                "description": t.get("description", ""),
                "service_module": t.get("service_module", ""),
                "severity": "urgent" if days_until < 0 else "warning" if days_until <= 3 else "info",
                "ref_id": t.get("task_id"),
            })

    # Upcoming events (within 14 days)
    events = await db.events.find({"date": {"$exists": True}}, {"_id": 0}).to_list(200)
    for ev in events:
        ev_date_str = ev.get("date")
        if not ev_date_str:
            continue
        try:
            ev_date = datetime.fromisoformat(ev_date_str).date() if 'T' in ev_date_str else datetime.strptime(ev_date_str, "%Y-%m-%d").date()
        except (ValueError, TypeError):
            continue
        days_until = (ev_date - today).days
        if 0 <= days_until <= 14:
            reminders.append({
                "reminder_id": f"rem_ev_{ev.get('event_id', '')}",
                "type": "event",
                "title": ev.get("title", "Event"),
                "client_name": ev.get("client_name", ""),
                "client_id": ev.get("client_id"),
                "assigned_to": ev.get("assigned_to_name") or ev.get("assigned_to", ""),
                "due_date": ev_date_str,
                "days_until": days_until,
                "priority": "High" if days_until <= 1 else "Medium",
                "status": "Scheduled",
                "description": ev.get("notes") or ev.get("description", ""),
                "service_module": ev.get("event_type", "event").capitalize(),
                "severity": "warning" if days_until <= 1 else "info",
                "ref_id": ev.get("event_id"),
            })

    # Active engagements nearing attention
    engs = await db.service_engagements.find({"status": "Active"}, {"_id": 0}).to_list(100)
    for eng in engs:
        total = sum(len(g.get("items", [])) for g in eng.get("checklist", []))
        done = sum(1 for g in eng.get("checklist", []) for item in g.get("items", []) if item.get("done"))
        progress = round(done / total * 100) if total > 0 else 0
        if progress < 50:
            reminders.append({
                "reminder_id": f"rem_eng_{eng.get('engagement_id', '')}",
                "type": "engagement",
                "title": eng.get("service_type", "").replace("_", " ").title(),
                "client_name": eng.get("client_name", ""),
                "client_id": eng.get("client_id"),
                "assigned_to": eng.get("assigned_to_name") or eng.get("assigned_to", ""),
                "due_date": eng.get("created_at", "")[:10],
                "days_until": None,
                "priority": "High" if progress < 25 else "Medium",
                "status": f"{progress}% complete — Phase: {eng.get('phase', '—')}",
                "description": f"Engagement at {progress}% completion, currently in {eng.get('phase', '—')} phase",
                "service_module": eng.get("service_type", "").replace("_", " ").title(),
                "severity": "warning" if progress < 25 else "info",
                "ref_id": eng.get("engagement_id"),
            })

    # Sort: urgent first
    severity_order = {"urgent": 0, "warning": 1, "info": 2}
    reminders.sort(key=lambda x: (severity_order.get(x["severity"], 2), x.get("due_date") or "9999"))
    return reminders

# ============= CLIENT PORTAL =============

# Predefined document checklists per service type
CLIENT_DOC_CHECKLISTS = {
    "aml": [
        {"label": "Emirates ID (Front & Back)", "required": True},
        {"label": "Passport Copy", "required": True},
        {"label": "Source of Funds Declaration", "required": True},
        {"label": "Beneficial Ownership Structure", "required": True},
        {"label": "Bank Statements (6 months)", "required": True},
        {"label": "Sanctions Self-Declaration", "required": True},
    ],
    "company_formation": [
        {"label": "Trade Licence Application Form", "required": True},
        {"label": "Passport Copies (All Shareholders)", "required": True},
        {"label": "Emirates ID Copies", "required": True},
        {"label": "Memorandum of Association (Draft)", "required": True},
        {"label": "NOC from Current Sponsor (if applicable)", "required": False},
        {"label": "Proof of Address", "required": True},
        {"label": "Business Plan", "required": False},
    ],
    "vat_registration": [
        {"label": "Trade Licence Copy", "required": True},
        {"label": "Passport / Emirates ID of Authorized Signatory", "required": True},
        {"label": "Bank Letter / IBAN Certificate", "required": True},
        {"label": "Turnover Evidence (12 months)", "required": True},
        {"label": "Import/Export Documentation (if applicable)", "required": False},
        {"label": "Lease Agreement / Ejari", "required": False},
    ],
    "audit": [
        {"label": "Trial Balance (Year End)", "required": True},
        {"label": "Financial Statements (Draft)", "required": True},
        {"label": "Bank Reconciliation Statements", "required": True},
        {"label": "Accounts Receivable Aging", "required": True},
        {"label": "Accounts Payable Aging", "required": True},
        {"label": "Fixed Asset Register", "required": True},
        {"label": "Inventory Listing", "required": False},
        {"label": "Payroll Summary", "required": False},
        {"label": "Related Party Transactions Detail", "required": False},
    ],
}

@api_router.get("/client-portal/documents")
async def get_client_documents(authorization: str = Header(None), session_token: str = Cookie(None)):
    """Get document checklist for the logged-in client user."""
    user = await get_current_user(authorization, session_token)
    if user.get("role") != "client":
        raise HTTPException(status_code=403, detail="Client role required")
    client_id = user.get("client_id")
    if not client_id:
        return {"checklist": [], "client_name": ""}

    client = await db.clients.find_one({"client_id": client_id}, {"_id": 0})
    client_name = client.get("name", "") if client else ""

    # Get or create document checklist for this client
    doc_checklist = await db.client_doc_checklists.find_one({"client_id": client_id}, {"_id": 0})
    if not doc_checklist:
        # Determine service type from client's active services or engagements
        service_type = None
        services = client.get("active_services", []) if client else []
        if any("AML" in s for s in services):
            service_type = "aml"
        elif any("Formation" in s for s in services):
            service_type = "company_formation"
        elif any("VAT Registration" in s for s in services):
            service_type = "vat_registration"
        elif any("Audit" in s for s in services):
            service_type = "audit"
        else:
            service_type = "audit"  # default

        template = CLIENT_DOC_CHECKLISTS.get(service_type, CLIENT_DOC_CHECKLISTS["audit"])
        items = [{"item_id": f"dci_{uuid.uuid4().hex[:8]}", "label": t["label"], "required": t["required"], "uploaded": False, "file_id": None, "filename": None, "uploaded_at": None} for t in template]

        doc_checklist = {
            "client_id": client_id,
            "service_type": service_type,
            "items": items,
            "created_at": datetime.now(timezone.utc).isoformat(),
        }
        await db.client_doc_checklists.insert_one(doc_checklist)
        doc_checklist.pop("_id", None)

    return {"checklist": doc_checklist.get("items", []), "service_type": doc_checklist.get("service_type"), "client_name": client_name}

@api_router.post("/client-portal/documents/{item_id}/upload")
async def upload_client_document(item_id: str, file: UploadFile = File(...), authorization: str = Header(None), session_token: str = Cookie(None)):
    """Client uploads a file for a specific checklist item."""
    user = await get_current_user(authorization, session_token)
    if user.get("role") != "client":
        raise HTTPException(status_code=403, detail="Client role required")
    client_id = user.get("client_id")
    if not client_id:
        raise HTTPException(status_code=403, detail="No client linked")

    # Verify item_id exists in this client's checklist
    checklist_doc = await db.client_doc_checklists.find_one({"client_id": client_id, "items.item_id": item_id}, {"_id": 0})
    if not checklist_doc:
        raise HTTPException(status_code=404, detail="Checklist item not found")

    content = await file.read()
    file_id = f"file_{uuid.uuid4().hex[:12]}"

    # Store file in documents collection
    doc_entry = {
        "file_id": file_id,
        "filename": file.filename,
        "content_type": file.content_type,
        "size": len(content),
        "data": content,
        "client_id": client_id,
        "uploader_id": user["user_id"],
        "uploader_name": user.get("name"),
        "uploaded_at": datetime.now(timezone.utc).isoformat(),
        "source": "client_portal",
    }
    await db.documents.insert_one(doc_entry)

    # Update checklist item
    await db.client_doc_checklists.update_one(
        {"client_id": client_id, "items.item_id": item_id},
        {"$set": {
            "items.$.uploaded": True,
            "items.$.file_id": file_id,
            "items.$.filename": file.filename,
            "items.$.uploaded_at": datetime.now(timezone.utc).isoformat(),
        }}
    )

    return {"file_id": file_id, "filename": file.filename, "message": "Document uploaded"}

@api_router.get("/client-portal/workflow")
async def get_client_workflow(authorization: str = Header(None), session_token: str = Cookie(None)):
    """Get workflow stages visible to client."""
    user = await get_current_user(authorization, session_token)
    if user.get("role") != "client":
        raise HTTPException(status_code=403, detail="Client role required")
    client_id = user.get("client_id")
    if not client_id:
        return {"engagements": []}

    engs = await db.service_engagements.find({"client_id": client_id}, {"_id": 0}).to_list(20)
    result = []
    for eng in engs:
        stages = []
        for group in eng.get("checklist", []):
            items = group.get("items", [])
            total = len(items)
            done = sum(1 for i in items if i.get("done"))
            status = "Completed" if done == total and total > 0 else "In Progress" if done > 0 else "Pending"
            stages.append({"name": group.get("group", ""), "status": status})

        result.append({
            "engagement_id": eng.get("engagement_id"),
            "service_type": eng.get("service_type", "").replace("_", " ").title(),
            "status": eng.get("status", "Active"),
            "phase": eng.get("phase", ""),
            "stages": stages,
        })

    return {"engagements": result}

@api_router.get("/client-portal/invoices")
async def get_client_invoices(authorization: str = Header(None), session_token: str = Cookie(None)):
    """Get invoice history for the logged-in client."""
    user = await get_current_user(authorization, session_token)
    if user.get("role") != "client":
        raise HTTPException(status_code=403, detail="Client role required")
    client_id = user.get("client_id")
    if not client_id:
        return []
    invoices = await db.invoices.find({"client_id": client_id}, {"_id": 0}).sort("date", -1).to_list(100)
    return invoices

# ============= INVOICES (ADMIN) =============

class InvoiceRequest(BaseModel):
    client_id: str
    client_name: Optional[str] = None
    service_name: str
    amount: float
    date: str
    status: str = "Unpaid"
    notes: Optional[str] = None

@api_router.post("/invoices")
async def create_invoice(req: InvoiceRequest, authorization: str = Header(None), session_token: str = Cookie(None)):
    user = await require_partner(authorization, session_token)
    if req.status not in ("Paid", "Unpaid"):
        raise HTTPException(status_code=400, detail="Status must be 'Paid' or 'Unpaid'")
    if req.amount < 0:
        raise HTTPException(status_code=400, detail="Amount must be non-negative")
    invoice_id = f"inv_{uuid.uuid4().hex[:12]}"
    client_name = req.client_name
    if not client_name:
        c = await db.clients.find_one({"client_id": req.client_id}, {"_id": 0, "name": 1})
        client_name = c["name"] if c else "Unknown"

    doc = {
        "invoice_id": invoice_id,
        "client_id": req.client_id,
        "client_name": client_name,
        "service_name": req.service_name,
        "amount": req.amount,
        "date": req.date,
        "status": req.status,
        "notes": req.notes,
        "created_by": user["user_id"],
        "created_at": datetime.now(timezone.utc).isoformat(),
    }
    await db.invoices.insert_one(doc)
    doc.pop("_id", None)
    return doc

@api_router.get("/invoices")
async def get_all_invoices(authorization: str = Header(None), session_token: str = Cookie(None), client_id: Optional[str] = None):
    user = await require_partner(authorization, session_token)
    query = {}
    if client_id:
        query["client_id"] = client_id
    invoices = await db.invoices.find(query, {"_id": 0}).sort("date", -1).to_list(500)
    return invoices

@api_router.patch("/invoices/{invoice_id}")
async def update_invoice(invoice_id: str, authorization: str = Header(None), session_token: str = Cookie(None), status: Optional[str] = None, amount: Optional[float] = None, notes: Optional[str] = None):
    user = await require_partner(authorization, session_token)
    inv = await db.invoices.find_one({"invoice_id": invoice_id})
    if not inv:
        raise HTTPException(status_code=404, detail="Invoice not found")
    update = {}
    if status:
        if status not in ("Paid", "Unpaid"):
            raise HTTPException(status_code=400, detail="Status must be 'Paid' or 'Unpaid'")
        update["status"] = status
    if amount is not None:
        update["amount"] = amount
    if notes is not None:
        update["notes"] = notes
    if update:
        await db.invoices.update_one({"invoice_id": invoice_id}, {"$set": update})
    return {"message": "Invoice updated"}

@api_router.delete("/invoices/{invoice_id}")
async def delete_invoice(invoice_id: str, authorization: str = Header(None), session_token: str = Cookie(None)):
    user = await require_partner(authorization, session_token)
    result = await db.invoices.delete_one({"invoice_id": invoice_id})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Invoice not found")
    return {"message": "Invoice deleted"}

# Admin: edit client document checklist
@api_router.get("/admin/client-checklists/{client_id}")
async def get_client_checklist(client_id: str, authorization: str = Header(None), session_token: str = Cookie(None)):
    user = await require_partner(authorization, session_token)
    doc = await db.client_doc_checklists.find_one({"client_id": client_id}, {"_id": 0})
    return doc or {"client_id": client_id, "items": []}

class ChecklistUpdateRequest(BaseModel):
    items: List[Dict[str, Any]]

@api_router.patch("/admin/client-checklists/{client_id}")
async def update_client_checklist(client_id: str, req: ChecklistUpdateRequest, authorization: str = Header(None), session_token: str = Cookie(None)):
    user = await require_partner(authorization, session_token)
    await db.client_doc_checklists.update_one(
        {"client_id": client_id},
        {"$set": {"items": req.items, "updated_by": user["user_id"], "updated_at": datetime.now(timezone.utc).isoformat()}},
        upsert=True
    )
    return {"message": "Checklist updated"}

# ============= WEB PUSH NOTIFICATIONS =============

@api_router.get("/push/vapid-key")
async def get_vapid_public_key():
    """Return the VAPID public key for frontend push subscription."""
    return {"public_key": VAPID_PUBLIC_KEY}

class PushSubscriptionRequest(BaseModel):
    subscription: Dict[str, Any]

@api_router.post("/push/subscribe")
async def subscribe_push(req: PushSubscriptionRequest, authorization: str = Header(None), session_token: str = Cookie(None)):
    """Store a push subscription for the authenticated user."""
    user = await get_current_user(authorization, session_token)
    endpoint = req.subscription.get("endpoint", "")

    # Upsert subscription (one per endpoint per user)
    await db.push_subscriptions.update_one(
        {"user_id": user["user_id"], "endpoint": endpoint},
        {"$set": {
            "user_id": user["user_id"],
            "user_email": user.get("email"),
            "user_name": user.get("name"),
            "subscription": req.subscription,
            "created_at": datetime.now(timezone.utc).isoformat(),
        }},
        upsert=True
    )
    return {"message": "Push subscription stored"}

@api_router.post("/push/unsubscribe")
async def unsubscribe_push(req: PushSubscriptionRequest, authorization: str = Header(None), session_token: str = Cookie(None)):
    """Remove a push subscription."""
    user = await get_current_user(authorization, session_token)
    endpoint = req.subscription.get("endpoint", "")
    await db.push_subscriptions.delete_one({"user_id": user["user_id"], "endpoint": endpoint})
    return {"message": "Push subscription removed"}

async def send_push_to_user(user_id: str, title: str, body: str, url: str = "/", tag: str = "nn-notif"):
    """Send a push notification to all subscriptions of a specific user."""
    from pywebpush import webpush, WebPushException

    if not VAPID_PRIVATE_KEY or not VAPID_PUBLIC_KEY:
        logger.warning("VAPID keys not configured, skipping push")
        return 0

    subs = await db.push_subscriptions.find({"user_id": user_id}, {"_id": 0}).to_list(20)
    sent = 0
    payload = json.dumps({
        "title": title,
        "body": body,
        "url": url,
        "tag": tag,
        "icon": "/nn-icon-192.png",
        "badge": "/nn-icon-192.png",
    })

    for sub_doc in subs:
        sub_info = sub_doc.get("subscription", {})
        try:
            webpush(
                subscription_info=sub_info,
                data=payload,
                vapid_private_key=VAPID_PRIVATE_KEY,
                vapid_claims={"sub": VAPID_CLAIMS_EMAIL},
            )
            sent += 1
        except WebPushException as e:
            logger.warning(f"Push failed for {sub_doc.get('user_email')}: {e}")
            # If subscription is expired/invalid (410 Gone), remove it
            if hasattr(e, 'response') and e.response is not None and e.response.status_code in (404, 410):
                endpoint = sub_info.get("endpoint", "")
                await db.push_subscriptions.delete_one({"user_id": user_id, "endpoint": endpoint})
        except Exception as e:
            logger.warning(f"Push error: {e}")

    return sent

@api_router.post("/push/test")
async def test_push(authorization: str = Header(None), session_token: str = Cookie(None)):
    """Send a test push notification to the current user."""
    user = await get_current_user(authorization, session_token)
    sent = await send_push_to_user(
        user["user_id"],
        "Nair & Nelliyatt",
        f"Push notifications are working, {user.get('name', 'there')}!",
        "/dashboard",
        "nn-test"
    )
    return {"message": f"Test push sent to {sent} device(s)"}

@api_router.post("/push/send-deadline-alerts")
async def send_deadline_push_alerts(authorization: str = Header(None), session_token: str = Cookie(None)):
    """Partner can trigger push alerts for overdue/due-soon tasks to all assigned staff."""
    user = await get_current_user(authorization, session_token)
    if user.get("role") != "partner":
        raise HTTPException(status_code=403, detail="Partners only")

    today = datetime.now(timezone.utc).date()
    tasks = await db.tasks.find({"status": {"$ne": "Completed"}}, {"_id": 0}).to_list(500)

    # Group alerts by assigned user email
    user_alerts = {}
    for t in tasks:
        due = t.get("due_date")
        if not due:
            continue
        try:
            due_date = datetime.fromisoformat(due).date() if 'T' in due else datetime.strptime(due, "%Y-%m-%d").date()
        except (ValueError, TypeError):
            continue
        days_until = (due_date - today).days
        if days_until > 3:
            continue

        assignee_email = t.get("assigned_to")
        if not assignee_email:
            continue

        # Find user_id for this email
        assignee = await db.users.find_one({"email": assignee_email}, {"_id": 0, "user_id": 1, "name": 1})
        if not assignee:
            continue

        uid = assignee["user_id"]
        if uid not in user_alerts:
            user_alerts[uid] = []
        label = f"OVERDUE: {t['title']}" if days_until < 0 else f"Due {'today' if days_until == 0 else f'in {days_until}d'}: {t['title']}"
        user_alerts[uid].append(label)

    total_sent = 0
    for uid, alerts in user_alerts.items():
        body = f"{len(alerts)} deadline alert{'s' if len(alerts) != 1 else ''}: {alerts[0]}" + (f" (+{len(alerts)-1} more)" if len(alerts) > 1 else "")
        sent = await send_push_to_user(uid, "Deadline Alert — N&N", body, "/tasks", "nn-deadline")
        total_sent += sent

    return {"message": f"Pushed alerts to {len(user_alerts)} user(s), {total_sent} device(s)"}

# ============= RESET DATA (ADMIN) =============

class ResetRequest(BaseModel):
    confirm: str
    collections: List[str] = []

RESETTABLE_COLLECTIONS = ["clients", "tasks", "events", "activities", "appreciations", "service_engagements", "documents", "invoices", "billable_hours", "client_doc_checklists", "vat_registrations", "vat_filings", "audit_engagements", "aml_alerts", "workflows", "chat_messages", "dismissed_notifications", "push_subscriptions"]

@api_router.post("/admin/reset-data")
async def reset_data(req: ResetRequest, authorization: str = Header(None), session_token: str = Cookie(None)):
    """Clear selected collections. Requires confirmation text 'RESET'."""
    user = await get_current_user(authorization, session_token)
    if user.get("role") != "partner" or user.get("title") != "Managing Partner":
        raise HTTPException(status_code=403, detail="Only Managing Partner can reset data")
    if req.confirm != "RESET":
        raise HTTPException(status_code=400, detail="Type 'RESET' to confirm")

    deleted_counts = {}
    targets = req.collections if req.collections else RESETTABLE_COLLECTIONS

    for coll_name in targets:
        if coll_name in RESETTABLE_COLLECTIONS:
            result = await db[coll_name].delete_many({})
            deleted_counts[coll_name] = result.deleted_count

    await log_activity("Data reset", f"Reset {len(deleted_counts)} collections: {', '.join(deleted_counts.keys())}", user["user_id"])
    return {"message": f"Reset {len(deleted_counts)} collection(s)", "deleted": deleted_counts}

@api_router.get("/admin/data-stats")
async def get_data_stats(authorization: str = Header(None), session_token: str = Cookie(None)):
    """Get document count per collection for the reset UI."""
    user = await get_current_user(authorization, session_token)
    if user.get("role") != "partner":
        raise HTTPException(status_code=403, detail="Partners only")

    stats = {}
    for coll_name in RESETTABLE_COLLECTIONS:
        stats[coll_name] = await db[coll_name].count_documents({})
    stats["users"] = await db.users.count_documents({})
    return stats

# ============= GOOGLE DRIVE INTEGRATION =============

from google_auth_oauthlib.flow import Flow
from google.oauth2.credentials import Credentials
from google.auth.transport.requests import Request as GoogleRequest
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseUpload

DRIVE_SCOPES = ['https://www.googleapis.com/auth/drive.file']

@api_router.get("/drive/connect")
async def connect_drive(authorization: str = Header(None), session_token: str = Cookie(None)):
    """Initiate Google Drive OAuth flow."""
    user = await require_partner(authorization, session_token)

    # Get Drive config from settings
    storage_doc = await db.settings.find_one({"type": "storage"}, {"_id": 0})
    config = storage_doc.get("config", {}).get("google_drive", {}) if storage_doc else {}

    client_id = config.get("client_id") or os.environ.get("GOOGLE_DRIVE_CLIENT_ID")
    client_secret = config.get("client_secret") or os.environ.get("GOOGLE_DRIVE_CLIENT_SECRET")

    if not client_id or not client_secret:
        raise HTTPException(status_code=400, detail="Google Drive credentials not configured. Go to Settings > Storage to add Client ID and Client Secret.")

    frontend_url = os.environ.get("FRONTEND_URL", "").rstrip("/")
    redirect_uri = f"{frontend_url}/api/drive/callback"

    flow = Flow.from_client_config(
        {"web": {"client_id": client_id, "client_secret": client_secret, "auth_uri": "https://accounts.google.com/o/oauth2/auth", "token_uri": "https://oauth2.googleapis.com/token", "redirect_uris": [redirect_uri]}},
        scopes=DRIVE_SCOPES,
        redirect_uri=redirect_uri
    )
    authorization_url, state = flow.authorization_url(access_type='offline', include_granted_scopes='true', prompt='consent', state=user["user_id"])
    return {"authorization_url": authorization_url}

@api_router.get("/drive/callback")
async def drive_callback(code: str = "", state: str = "", error: str = ""):
    """Handle Google Drive OAuth callback."""
    if error:
        raise HTTPException(status_code=400, detail=f"OAuth error: {error}")

    storage_doc = await db.settings.find_one({"type": "storage"}, {"_id": 0})
    config = storage_doc.get("config", {}).get("google_drive", {}) if storage_doc else {}
    client_id = config.get("client_id") or os.environ.get("GOOGLE_DRIVE_CLIENT_ID")
    client_secret = config.get("client_secret") or os.environ.get("GOOGLE_DRIVE_CLIENT_SECRET")

    frontend_url = os.environ.get("FRONTEND_URL", "").rstrip("/")
    redirect_uri = f"{frontend_url}/api/drive/callback"

    flow = Flow.from_client_config(
        {"web": {"client_id": client_id, "client_secret": client_secret, "auth_uri": "https://accounts.google.com/o/oauth2/auth", "token_uri": "https://oauth2.googleapis.com/token", "redirect_uris": [redirect_uri]}},
        scopes=None,
        redirect_uri=redirect_uri
    )
    flow.fetch_token(code=code)
    credentials = flow.credentials

    await db.drive_credentials.update_one(
        {"type": "firm"},
        {"$set": {
            "type": "firm",
            "connected_by": state,
            "access_token": credentials.token,
            "refresh_token": credentials.refresh_token,
            "token_uri": credentials.token_uri,
            "client_id": credentials.client_id,
            "client_secret": credentials.client_secret,
            "scopes": list(credentials.scopes) if credentials.scopes else [],
            "expiry": credentials.expiry.isoformat() if credentials.expiry else None,
            "connected_at": datetime.now(timezone.utc).isoformat(),
        }},
        upsert=True
    )

    from starlette.responses import RedirectResponse
    return RedirectResponse(url=f"{frontend_url}/app/settings?drive=connected")

@api_router.get("/drive/status")
async def drive_status(authorization: str = Header(None), session_token: str = Cookie(None)):
    """Check if Google Drive is connected."""
    user = await get_current_user(authorization, session_token)
    creds_doc = await db.drive_credentials.find_one({"type": "firm"}, {"_id": 0})
    if not creds_doc or not creds_doc.get("access_token"):
        return {"connected": False}
    return {"connected": True, "connected_at": creds_doc.get("connected_at"), "connected_by": creds_doc.get("connected_by")}

async def get_drive_service():
    """Get an authenticated Drive service with auto-refresh."""
    creds_doc = await db.drive_credentials.find_one({"type": "firm"}, {"_id": 0})
    if not creds_doc or not creds_doc.get("access_token"):
        return None

    creds = Credentials(
        token=creds_doc["access_token"],
        refresh_token=creds_doc.get("refresh_token"),
        token_uri=creds_doc["token_uri"],
        client_id=creds_doc["client_id"],
        client_secret=creds_doc["client_secret"],
        scopes=creds_doc.get("scopes"),
    )

    if creds.expired and creds.refresh_token:
        creds.refresh(GoogleRequest())
        await db.drive_credentials.update_one(
            {"type": "firm"},
            {"$set": {"access_token": creds.token, "expiry": creds.expiry.isoformat() if creds.expiry else None}}
        )

    return build('drive', 'v3', credentials=creds)

@api_router.post("/drive/sync-document/{file_id}")
async def sync_document_to_drive(file_id: str, authorization: str = Header(None), session_token: str = Cookie(None)):
    """Upload a specific document to Google Drive."""
    user = await require_partner(authorization, session_token)
    service = await get_drive_service()
    if not service:
        raise HTTPException(status_code=400, detail="Google Drive not connected")

    file_doc = await db.documents.find_one({"file_id": file_id}, {"_id": 0})
    if not file_doc:
        raise HTTPException(status_code=404, detail="File not found")

    # Get or create a folder for the client
    client_name = file_doc.get("client_name") or file_doc.get("client_id", "General")
    folder_id = await get_or_create_drive_folder(service, client_name)

    # Upload file
    file_data = file_doc.get("data")
    if not file_data:
        raise HTTPException(status_code=400, detail="No file data to sync")

    media = MediaIoBaseUpload(io.BytesIO(file_data), mimetype=file_doc.get("content_type", "application/octet-stream"))
    file_metadata = {"name": file_doc.get("filename", "document"), "parents": [folder_id]}
    drive_file = service.files().create(body=file_metadata, media_body=media, fields="id,name,webViewLink").execute()

    # Update document record with drive link
    await db.documents.update_one(
        {"file_id": file_id},
        {"$set": {"drive_file_id": drive_file["id"], "drive_link": drive_file.get("webViewLink"), "synced_at": datetime.now(timezone.utc).isoformat()}}
    )

    return {"drive_file_id": drive_file["id"], "drive_link": drive_file.get("webViewLink"), "message": "Synced to Google Drive"}

async def get_or_create_drive_folder(service, folder_name):
    """Get or create a folder in Drive for organizing client docs."""
    # First check if we have a root NN folder
    root_query = "name='Nair & Nelliyatt Documents' and mimeType='application/vnd.google-apps.folder' and trashed=false"
    results = service.files().list(q=root_query, fields="files(id,name)", spaces='drive').execute()
    root_files = results.get("files", [])

    if root_files:
        root_id = root_files[0]["id"]
    else:
        root_meta = {"name": "Nair & Nelliyatt Documents", "mimeType": "application/vnd.google-apps.folder"}
        root_folder = service.files().create(body=root_meta, fields="id").execute()
        root_id = root_folder["id"]

    # Now get or create client subfolder
    client_query = f"name='{folder_name}' and mimeType='application/vnd.google-apps.folder' and '{root_id}' in parents and trashed=false"
    results = service.files().list(q=client_query, fields="files(id,name)", spaces='drive').execute()
    client_files = results.get("files", [])

    if client_files:
        return client_files[0]["id"]
    else:
        client_meta = {"name": folder_name, "mimeType": "application/vnd.google-apps.folder", "parents": [root_id]}
        client_folder = service.files().create(body=client_meta, fields="id").execute()
        return client_folder["id"]

@api_router.get("/drive/files")
async def list_drive_files(authorization: str = Header(None), session_token: str = Cookie(None)):
    """List files in the N&N Drive folder."""
    user = await require_partner(authorization, session_token)
    service = await get_drive_service()
    if not service:
        raise HTTPException(status_code=400, detail="Google Drive not connected")

    root_query = "name='Nair & Nelliyatt Documents' and mimeType='application/vnd.google-apps.folder' and trashed=false"
    results = service.files().list(q=root_query, fields="files(id,name)", spaces='drive').execute()
    root_files = results.get("files", [])
    if not root_files:
        return {"files": [], "folder_count": 0}

    root_id = root_files[0]["id"]
    # List subfolders
    folders_query = f"'{root_id}' in parents and mimeType='application/vnd.google-apps.folder' and trashed=false"
    folders = service.files().list(q=folders_query, fields="files(id,name)", spaces='drive').execute().get("files", [])

    return {"root_folder_id": root_id, "folders": folders, "folder_count": len(folders)}

@api_router.post("/drive/disconnect")
async def disconnect_drive(authorization: str = Header(None), session_token: str = Cookie(None)):
    """Disconnect Google Drive."""
    user = await require_partner(authorization, session_token)
    await db.drive_credentials.delete_many({"type": "firm"})
    return {"message": "Google Drive disconnected"}

# Include router
app.include_router(api_router)

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=os.environ.get('CORS_ORIGINS', '*').split(','),
    allow_methods=["*"],
    allow_headers=["*"],
)

async def seed_nn_users():
    """Seed the Nair & Nelliyatt team into the database if not already present."""
    default_password = hash_password("nn123456")
    team = [
        {"name": "Arjun Srinivas", "email": "arjun@nnadvisory.ae", "role": "partner", "title": "Managing Partner"},
        {"name": "Sooraj Nelliyatt", "email": "sooraj@nnadvisory.ae", "role": "partner", "title": "Senior Partner"},
        {"name": "Fazil", "email": "fazil@nnadvisory.ae", "role": "staff", "title": "Associate"},
        {"name": "Subin", "email": "subin@nnadvisory.ae", "role": "staff", "title": "Associate"},
        {"name": "Anju", "email": "anju@nnadvisory.ae", "role": "staff", "title": "Senior Associate"},
        {"name": "Roshith", "email": "roshith@nnadvisory.ae", "role": "staff", "title": "Associate"},
        {"name": "Thasleema", "email": "thasleema@nnadvisory.ae", "role": "staff", "title": "Senior Associate"},
        {"name": "Jithin", "email": "jithin@nnadvisory.ae", "role": "staff", "title": "Associate"},
        {"name": "Shamil A.", "email": "shamil@nnadvisory.ae", "role": "staff", "title": "Associate"},
        {"name": "Akhil", "email": "akhil@nnadvisory.ae", "role": "staff", "title": "Associate"},
        {"name": "Haritha", "email": "haritha@nnadvisory.ae", "role": "staff", "title": "Senior Associate"},
    ]
    for member in team:
        existing = await db.users.find_one({"email": member["email"]})
        if not existing:
            await db.users.insert_one({
                "user_id": f"user_{uuid.uuid4().hex[:12]}",
                "email": member["email"],
                "name": member["name"],
                "title": member["title"],
                "role": member["role"],
                "password": default_password,
                "created_at": datetime.now(timezone.utc).isoformat()
            })
            logger.info(f"Seeded user: {member['name']}")

async def seed_nn_clients():
    """Seed demo clients for Nair & Nelliyatt practice."""
    demo_clients = [
        {"name": "Al Baraka Trading LLC", "entity_type": "LLC", "jurisdiction": "Dubai Mainland", "trn": "100234567890003", "status": "Active", "active_services": ["Statutory Audit", "VAT Filing", "AML"]},
        {"name": "Falcon Logistics Co.", "entity_type": "LLC", "jurisdiction": "Dubai Mainland", "trn": "100456789000123", "status": "Active", "active_services": ["Statutory Audit", "VAT Filing", "Internal Audit", "AML"]},
        {"name": "Gulf Pharma Group", "entity_type": "Group", "jurisdiction": "Abu Dhabi", "trn": "100345678900012", "status": "Active", "active_services": ["Statutory Audit", "VAT Filing", "Corporate Tax", "AML"]},
        {"name": "Zara Tech LLC", "entity_type": "LLC", "jurisdiction": "DMCC Free Zone", "trn": "100678900012345", "status": "Active", "active_services": ["Statutory Audit", "VAT Filing"]},
        {"name": "Sunrise Holdings", "entity_type": "Holding", "jurisdiction": "Dubai Mainland", "status": "Active", "active_services": ["Statutory Audit", "Company Formation"]},
        {"name": "Marina Holdings", "entity_type": "LLC", "jurisdiction": "Dubai Marina", "trn": "100567890001234", "status": "Active", "active_services": ["Statutory Audit", "VAT Filing", "Internal Audit", "Corporate Tax"]},
        {"name": "Desert Rose Trading", "entity_type": "LLC", "jurisdiction": "Sharjah", "status": "Active", "active_services": ["Statutory Audit", "AML", "Valuation"]},
        {"name": "Al Hayat Retail", "entity_type": "LLC", "jurisdiction": "Dubai Mainland", "status": "Active", "active_services": ["Statutory Audit"]},
        {"name": "Blue Horizon Co.", "entity_type": "LLC", "jurisdiction": "Ajman Free Zone", "status": "Active", "active_services": ["Statutory Audit"]},
        {"name": "Bright Vision LLC", "entity_type": "LLC", "jurisdiction": "Dubai Mainland", "status": "Active", "active_services": ["Statutory Audit"]},
    ]
    for c in demo_clients:
        existing = await db.clients.find_one({"name": c["name"]})
        if not existing:
            await db.clients.insert_one({
                "client_id": f"client_{uuid.uuid4().hex[:12]}",
                "name": c["name"],
                "entity_type": c["entity_type"],
                "jurisdiction": c.get("jurisdiction"),
                "trn": c.get("trn"),
                "status": c.get("status", "Active"),
                "active_services": c.get("active_services", []),
                "created_at": datetime.now(timezone.utc).isoformat()
            })
            logger.info(f"Seeded client: {c['name']}")

async def seed_nn_engagements():
    """Seed sample service engagements if none exist."""
    existing = await db.service_engagements.count_documents({})
    if existing > 0:
        return

    clients = await db.clients.find({}, {"_id": 0, "client_id": 1, "name": 1}).to_list(10)
    staff = await db.users.find({"role": "staff"}, {"_id": 0, "name": 1, "email": 1}).to_list(10)
    if not clients or not staff:
        return

    def make_checklist(svc_type, done_map):
        template = CHECKLIST_TEMPLATES.get(svc_type, [])
        cl = []
        for g in template:
            items = []
            for item in g["items"]:
                items.append({"label": item, "done": done_map.get(item, False)})
            cl.append({"group": g["group"], "items": items})
        return cl

    seed_engs = [
        {"service_type": "statutory_audit", "client_idx": 0, "staff_idx": 0, "phase": "Fieldwork",
         "done": {"Engagement letter": True, "Risk assessment": True, "Audit plan": True, "Team briefing": True, "Revenue testing": True}},
        {"service_type": "statutory_audit", "client_idx": 1, "staff_idx": 3, "phase": "Planning",
         "done": {"Engagement letter": True, "Risk assessment": True}},
        {"service_type": "statutory_audit", "client_idx": 2, "staff_idx": 4, "phase": "Fieldwork",
         "done": {"Engagement letter": True, "Risk assessment": True, "Audit plan": True, "Materiality memo": True, "Team briefing": True, "Revenue testing": True, "Expense sampling": True, "Bank confirmations": True}},
        {"service_type": "statutory_audit", "client_idx": 7, "staff_idx": 6, "phase": "Review & Reporting",
         "done": {"Engagement letter": True, "Risk assessment": True, "Audit plan": True, "Materiality memo": True, "Team briefing": True, "Revenue testing": True, "Expense sampling": True, "Bank confirmations": True, "Inventory count": True, "Related party review": True, "Draft report": True, "Management letter": True}},
        {"service_type": "vat_filing", "client_idx": 0, "staff_idx": 3, "phase": "Completion",
         "done": {"Data collection": True, "Purchase invoices review": True, "Sales invoices review": True, "Reconciliation": True, "Return preparation": True, "Box amounts calculation": True, "Review & approval": True, "FTA portal submission": True, "Payment confirmation": True, "Filing receipt archive": True, "Client notification": True}},
        {"service_type": "vat_filing", "client_idx": 2, "staff_idx": 4, "phase": "Completion",
         "done": {"Data collection": True, "Purchase invoices review": True, "Sales invoices review": True, "Reconciliation": True, "Return preparation": True, "Box amounts calculation": True, "Review & approval": True, "FTA portal submission": True, "Payment confirmation": True}},
        {"service_type": "vat_filing", "client_idx": 1, "staff_idx": 5, "phase": "Filing",
         "done": {"Data collection": True, "Purchase invoices review": True, "Sales invoices review": True, "Reconciliation": True, "Return preparation": True}},
        {"service_type": "vat_filing", "client_idx": 5, "staff_idx": 6, "phase": "Completion",
         "done": {"Data collection": True, "Purchase invoices review": True, "Sales invoices review": True, "Reconciliation": True, "Return preparation": True, "Box amounts calculation": True, "Review & approval": True, "FTA portal submission": True, "Payment confirmation": True, "Filing receipt archive": True}},
        {"service_type": "aml_review", "client_idx": 0, "staff_idx": 4, "phase": "Transaction Monitoring",
         "done": {"KYC documentation": True, "Beneficial ownership check": True, "PEP screening": True, "Sanctions screening": True, "Unusual transaction review": True}},
        {"service_type": "aml_review", "client_idx": 1, "staff_idx": 6, "phase": "Client Due Diligence",
         "done": {"KYC documentation": True, "Beneficial ownership check": True}},
        {"service_type": "corporate_tax", "client_idx": 2, "staff_idx": 4, "phase": "Computation",
         "done": {"Financial data collection": True, "Revenue classification": True, "Exempt income review": True, "Deduction analysis": True, "Taxable income calculation": True}},
    ]

    for se in seed_engs:
        ci = min(se["client_idx"], len(clients) - 1)
        si = min(se["staff_idx"], len(staff) - 1)
        checklist = make_checklist(se["service_type"], se.get("done", {}))
        total = sum(len(g["items"]) for g in checklist)
        done_count = sum(1 for g in checklist for item in g["items"] if item["done"])
        await db.service_engagements.insert_one({
            "engagement_id": f"eng_{uuid.uuid4().hex[:12]}",
            "service_type": se["service_type"],
            "client_id": clients[ci]["client_id"],
            "client_name": clients[ci]["name"],
            "assigned_to": staff[si]["email"],
            "assigned_to_name": staff[si]["name"],
            "status": "Completed" if done_count == total else "Active",
            "phase": se["phase"],
            "checklist": checklist,
            "created_by": "system",
            "created_at": datetime.now(timezone.utc).isoformat(),
        })
    logger.info(f"Seeded {len(seed_engs)} service engagements")

async def seed_preset_workflows():
    """Seed preset workflows for each service type if none exist."""
    existing = await db.workflows.count_documents({})
    if existing > 0:
        return
    presets = [
        {"name": "Statutory Audit Workflow", "service_type": "statutory_audit", "steps": [
            {"name": "Engagement Letter", "description": "Issue and sign engagement letter with client", "order": 0},
            {"name": "Risk Assessment", "description": "Assess inherent and control risks", "order": 1},
            {"name": "Audit Planning", "description": "Develop audit plan and materiality thresholds", "order": 2},
            {"name": "Fieldwork Execution", "description": "Perform substantive testing and controls testing", "order": 3},
            {"name": "Review & Quality Control", "description": "Partner review of working papers", "order": 4},
            {"name": "Draft Report", "description": "Prepare draft audit report and management letter", "order": 5},
            {"name": "Client Sign-off", "description": "Obtain client approval and representations", "order": 6},
            {"name": "Final Report & Filing", "description": "Issue final report and file with authorities", "order": 7},
        ]},
        {"name": "VAT Filing Workflow", "service_type": "vat_filing", "steps": [
            {"name": "Data Collection", "description": "Gather sales and purchase invoices from client", "order": 0},
            {"name": "Invoice Reconciliation", "description": "Reconcile invoices with accounting records", "order": 1},
            {"name": "Return Preparation", "description": "Calculate box amounts and prepare VAT return", "order": 2},
            {"name": "Partner Review", "description": "Review and approve return before submission", "order": 3},
            {"name": "FTA Submission", "description": "Submit return via FTA portal", "order": 4},
            {"name": "Payment Processing", "description": "Process VAT payment and confirm receipt", "order": 5},
        ]},
        {"name": "AML Review Workflow", "service_type": "aml_review", "steps": [
            {"name": "KYC Documentation", "description": "Collect and verify identity documents", "order": 0},
            {"name": "Beneficial Ownership", "description": "Identify and verify beneficial owners", "order": 1},
            {"name": "PEP & Sanctions Screening", "description": "Screen against PEP lists and sanctions databases", "order": 2},
            {"name": "Risk Rating", "description": "Assign client risk rating based on due diligence", "order": 3},
            {"name": "Transaction Monitoring", "description": "Review transactions for unusual patterns", "order": 4},
            {"name": "MLRO Review & goAML", "description": "MLRO assessment and goAML report if required", "order": 5},
        ]},
        {"name": "Corporate Tax Workflow", "service_type": "corporate_tax", "steps": [
            {"name": "Financial Data Review", "description": "Review financial statements and trial balance", "order": 0},
            {"name": "Income Classification", "description": "Classify taxable, exempt, and qualifying income", "order": 1},
            {"name": "Tax Computation", "description": "Calculate taxable income and tax liability", "order": 2},
            {"name": "Return Preparation", "description": "Prepare CT return with supporting schedules", "order": 3},
            {"name": "Partner Approval", "description": "Review and sign-off by engagement partner", "order": 4},
            {"name": "Filing & Payment", "description": "Submit via portal and process payment", "order": 5},
        ]},
        {"name": "Internal Audit Workflow", "service_type": "internal_audit", "steps": [
            {"name": "Scope Definition", "description": "Define audit scope and objectives", "order": 0},
            {"name": "Risk Assessment", "description": "Update risk universe and prioritize areas", "order": 1},
            {"name": "Audit Program", "description": "Develop detailed testing procedures", "order": 2},
            {"name": "Control Testing", "description": "Test operating effectiveness of controls", "order": 3},
            {"name": "Findings & Recommendations", "description": "Document findings and propose improvements", "order": 4},
            {"name": "Management Response", "description": "Obtain management responses to findings", "order": 5},
            {"name": "Final Report", "description": "Issue final internal audit report", "order": 6},
        ]},
    ]
    for p in presets:
        await db.workflows.insert_one({
            "workflow_id": f"wf_{uuid.uuid4().hex[:12]}",
            "name": p["name"],
            "service_type": p["service_type"],
            "steps": p["steps"],
            "is_preset": True,
            "created_by": "system",
            "created_at": datetime.now(timezone.utc).isoformat(),
        })
    logger.info(f"Seeded {len(presets)} preset workflows")

@app.on_event("startup")
async def startup():
    try:
        init_storage()
        logger.info("Storage initialized")
    except Exception as e:
        logger.error(f"Storage init failed: {e}")
    await seed_nn_users()
    await seed_nn_clients()
    await seed_nn_engagements()
    await seed_preset_workflows()

@app.on_event("shutdown")
async def shutdown_db_client():
    client.close()
